"""Steps, key figures, a pyramid of levels and a set of symbols, laid out for whatever count they hold."""

from __future__ import annotations

import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from . import colors as fb
from . import icons, grid
from .elements import _SLACK, MIN_PT, LINE_RATIO, font_for, style_run
from .colors import Palette
from . import shapes
from .shapes import _STEP, _BLACK, _emu, _fill, _shape, _textbox
from .schemas import Entry, Figure, Symbol
from .template import EMU_PER_PT, Field

# A step's circle in multiples of the base size, at most this share of its column, its number inside.
_CIRCLE = 1.2
_CIRCLE_SHARE = 0.5
_NUMBER = 0.45
# The line joining the steps, as a share of the circle, and the air between steps below each other.
_JOIN = 0.08
_ROW_AIR = 0.6
# A key figure at most this many times the base size, the lines between them this thin.
_MAX_VALUE_PT = 3.6
_DIVIDER_PT = 1.5
# A pyramid is about as tall as an equilateral one of its base, the lines to its texts this thin.
_PYRAMID_HEIGHT = 0.87
_CONNECTOR_PT = 1.0
# A symbol's circle in multiples of the base size, at most this share of its column, the icon this share of it.
_SYMBOL_CIRCLE = 4.0
_SYMBOL_CIRCLE_SHARE = 0.6
_ICON_IN_CIRCLE = 0.52
# A drawing sits a third of the free height below its field's top, closer to the heading than to the foot.
_TOP_SHARE = 0.35
# A cited provision breaks as a whole, a section sign at least keeps its number.
_BRACKET = re.compile(r"\((?:§|Art\.)[^()]*\)")
_REFERENCE = re.compile(r"(§§?|Art\.)\s+(?=\d)")
_NBSP = "\u00a0"
# Measured in place of the protected space, a letter as wide as a space that no word wrap splits at.
_AS_WORD = "i"


def _metrics(field: Field) -> tuple[float, float, float, float]:
    return field.left / EMU_PER_PT, field.top / EMU_PER_PT, field.width / EMU_PER_PT, field.height / EMU_PER_PT


def _bound(text: str) -> str:
    text = _BRACKET.sub(lambda m: m.group(0).replace(" ", _NBSP), text)
    return _REFERENCE.sub(lambda m: m.group(1) + _NBSP, text)


def _lines(text: str, typeface: str, size: float, width: float, bold: bool = False) -> int:
    return shapes._lines(_bound(text).replace(_NBSP, _AS_WORD), typeface, size, width, bold)


def _text(slide, x: float, y: float, w: float, h: float, text: str, size: float, color, bold: bool,
          typeface: str, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP):
    tf = _textbox(slide, x, y, w, h, anchor, align)
    run = tf.paragraphs[0].add_run()
    run.text = _bound(text)
    style_run(run, size, color, bold, typeface)
    return tf


def _whole(texts: list[str], font, width: float) -> bool:
    """No word of these texts is wider than the line, so none gets broken."""
    return all(font.getlength(w) <= width * _SLACK
               for t in texts for w in _bound(t).replace(_NBSP, _AS_WORD).split())


def _largest(measure, g: float, ratio: float = grid.TEXT_RATIO) -> tuple[bool, float, tuple]:
    """Whether it fits, the scale and the measures: the largest scale that fits, else the smallest readable one.

    The scale starts where the running text is as large as the template's body text.
    """
    k = 1 / ratio
    while True:
        fits, mass = measure(k)
        if fits or ratio * g * (k - _STEP) < MIN_PT:
            return fits, k, mass
        k -= _STEP


def _better(a: tuple[bool, float, tuple], b: tuple[bool, float, tuple] | None) -> bool:
    """A layout that fits beats one that does not, among those the larger text wins, a tie keeps the earlier one."""
    return b is None or (a[0], a[1]) > (b[0], b[1] + 1e-9)


def _entry_height(e: Entry, typeface: str, title_pt: float, text_pt: float, width: float) -> float:
    h = _lines(e.title, typeface, title_pt, width, True) * title_pt * LINE_RATIO
    if e.text:
        h += 0.3 * text_pt + _lines(e.text, typeface, text_pt, width) * text_pt * LINE_RATIO
    return h


def _entry(slide, x: float, y: float, width: float, e: Entry, title_pt: float, text_pt: float, f: Palette,
             typeface: str, align=PP_ALIGN.CENTER) -> None:
    """Title in the colour and bold, the sentence below it in black."""
    title_h = _lines(e.title, typeface, title_pt, width, True) * title_pt * LINE_RATIO
    _text(slide, x, y, width, title_h, e.title, title_pt, f.text_color(), True, typeface, align)
    if e.text:
        y += title_h + 0.3 * text_pt
        h = _lines(e.text, typeface, text_pt, width) * text_pt * LINE_RATIO
        _text(slide, x, y, width, h, e.text, text_pt, _BLACK, False, typeface, align)


# Steps


def steps_figure(slide, field: Field, steps: list[Entry], f: Palette, typeface: str) -> None:
    """Numbered circles on a line in reading order, side by side, or one below the other when that keeps the text larger."""
    typeface = typeface or "Calibri"
    g = grid.base_pt(field.size_pt)
    left, top, width, height = _metrics(field)
    n = len(steps)
    column = width / n
    inner = column - g

    def whole(title_pt: float, text_pt: float, w: float) -> bool:
        return (_whole([s.title for s in steps], font_for(typeface, title_pt, True), w)
                and _whole([s.text for s in steps], font_for(typeface, text_pt), w))

    def across(k: float):
        title_pt, text_pt = grid.TITLE_RATIO * g * k, grid.TEXT_RATIO * g * k
        d = min(_CIRCLE * g * k, column * _CIRCLE_SHARE)
        h = d + 0.5 * title_pt + max(_entry_height(s, typeface, title_pt, text_pt, inner) for s in steps)
        return whole(title_pt, text_pt, inner) and h <= height, (title_pt, text_pt, d, h, [], 0.0)

    def stacked_fit(k: float):
        title_pt, text_pt = grid.TITLE_RATIO * g * k, grid.TEXT_RATIO * g * k
        d = _CIRCLE * g * k
        text_w = width - d - g
        lines = [max(d, _entry_height(s, typeface, title_pt, text_pt, text_w)) for s in steps]
        air = _ROW_AIR * g * k
        h = sum(lines) + (n - 1) * air
        return whole(title_pt, text_pt, text_w) and h <= height, (title_pt, text_pt, d, h, lines, air)

    side_by_side = _largest(across, g)
    stacked = _largest(stacked_fit, g)
    vertical = n > 1 and _better(stacked, side_by_side)
    title_pt, text_pt, d, h, lines, air = (stacked if vertical else side_by_side)[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    thickness = max(1.5, d * _JOIN)
    if vertical:
        tops = [top + sum(lines[:i]) + i * air for i in range(n)]
        _shape(slide, MSO_SHAPE.RECTANGLE, left + (d - thickness) / 2, tops[0] + d / 2, thickness, tops[-1] - tops[0],
              f.tint(fb.LINE))
        circles = [(left, y) for y in tops]
    else:
        centres = [left + column * (i + 0.5) for i in range(n)]
        if n > 1:
            _shape(slide, MSO_SHAPE.RECTANGLE, centres[0], top + (d - thickness) / 2, centres[-1] - centres[0], thickness,
                  f.tint(fb.LINE))
        circles = [(x - d / 2, top) for x in centres]
    for i, (s, (x, y)) in enumerate(zip(steps, circles)):
        _shape(slide, MSO_SHAPE.OVAL, x, y, d, d, f.tint(fb.FULL))
        tf = _text(slide, x, y, d, d, str(i + 1), d * _NUMBER, f.text_on(), True, typeface,
                   anchor=MSO_ANCHOR.MIDDLE)
        tf.word_wrap = False
        if vertical:
            text_w = width - d - g
            block = _entry_height(s, typeface, title_pt, text_pt, text_w)
            _entry(slide, x + d + g, y + max(0.0, (d - block) / 2), text_w, s, title_pt, text_pt, f, typeface,
                     PP_ALIGN.LEFT)
        else:
            _entry(slide, x + d / 2 - inner / 2, y + d + 0.5 * title_pt, inner, s, title_pt, text_pt, f, typeface)


# Key figures


def figures(slide, field: Field, numbers: list[Figure], f: Palette, typeface: str) -> None:
    """Each figure large in the colour with its label below, thin lines between the columns."""
    typeface = typeface or "Calibri"
    g = grid.base_pt(field.size_pt)
    left, top, width, height = _metrics(field)
    column = width / len(numbers)
    inner = column - 1.5 * g
    # One size for every figure, the largest at which the widest one stays on one line.
    number_pt = min(_MAX_VALUE_PT * g, height * 0.4)
    while number_pt > MIN_PT and max(font_for(typeface, number_pt, True).getlength(z.value)
                                   for z in numbers) > inner * _SLACK:
        number_pt -= 1

    def measure(k: float):
        text_pt = g * k
        lines = max(_lines(z.label, typeface, text_pt, inner) for z in numbers)
        h = number_pt * LINE_RATIO + 0.4 * text_pt + lines * text_pt * LINE_RATIO
        return _whole([z.label for z in numbers], font_for(typeface, text_pt), inner) and h <= height, (text_pt, h)

    text_pt, h = _largest(measure, g, 1.0)[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    for i, z in enumerate(numbers):
        x = left + i * column
        if i:
            _shape(slide, MSO_SHAPE.RECTANGLE, x - _DIVIDER_PT / 2, top, _DIVIDER_PT, h, f.tint(fb.LINE))
        tf = _text(slide, x + (column - inner) / 2, top, inner, number_pt * LINE_RATIO, z.value, number_pt, f.text_color(),
                   True, typeface)
        tf.word_wrap = False
        y = top + number_pt * LINE_RATIO + 0.4 * text_pt
        _text(slide, x + (column - inner) / 2, y, inner, top + h - y, z.label, text_pt, _BLACK, False, typeface)


# Levels


def _polygon(slide, points: list[tuple[float, float]], color):
    """A closed polygon from points on the slide, the same outline in PowerPoint and LibreOffice."""
    x0, y0 = points[0]
    builder = slide.shapes.build_freeform(_emu(x0), _emu(y0), scale=1.0)
    builder.add_line_segments([(_emu(x), _emu(y)) for x, y in points[1:]], close=True)
    return _fill(builder.convert_to_shape(), color)


def levels(slide, field: Field, levels_: list[Entry], f: Palette, typeface: str) -> None:
    """A pyramid in tints of the colour, the top level in full colour, each level's title and sentence beside it."""
    typeface = typeface or "Calibri"
    g = grid.base_pt(field.size_pt)
    left, top, width, height = _metrics(field)
    n = len(levels_)
    gap = grid.LEVEL_GAP * g
    base = width * grid.PYRAMID_RATIO
    text_x = left + base + 1.5 * g
    text_w = left + width - text_x

    def measure(k: float):
        title_pt, text_pt = grid.TITLE_RATIO * g * k, grid.TEXT_RATIO * g * k
        whole = (_whole([e.title for e in levels_], font_for(typeface, title_pt, True), text_w)
                and _whole([e.text for e in levels_], font_for(typeface, text_pt), text_w))
        level = max(grid.LEVEL_MIN_HEIGHT * g * k,
                    max(_entry_height(e, typeface, title_pt, text_pt, text_w) for e in levels_))
        return whole and n * level + (n - 1) * gap <= height, (title_pt, text_pt, level)

    title_pt, text_pt, level = _largest(measure, g)[2]
    # The levels grow until the pyramid has its natural height, as far as the field allows.
    level = max(level, min((height - (n - 1) * gap) / n, (base * _PYRAMID_HEIGHT - (n - 1) * gap) / n))
    h = n * level + (n - 1) * gap
    top += max(0.0, (height - h) * _TOP_SHARE)
    apex = left + base / 2

    def half(y: float) -> float:
        return (y - top) / h * base / 2

    for i, (e, color) in enumerate(zip(levels_, f.levels(n))):
        ya = top + i * (level + gap)
        yb = ya + level
        head = [(apex, ya)] if i == 0 else [(apex - half(ya), ya), (apex + half(ya), ya)]
        _polygon(slide, head + [(apex + half(yb), yb), (apex - half(yb), yb)], color)
        middle = ya + level / 2
        pad = apex + half(middle) + 0.4 * g
        _shape(slide, MSO_SHAPE.RECTANGLE, pad, middle - _CONNECTOR_PT / 2, text_x - 0.4 * g - pad, _CONNECTOR_PT,
              f.tint(fb.LINE))
        block = _entry_height(e, typeface, title_pt, text_pt, text_w)
        _entry(slide, text_x, middle - block / 2, text_w, e, title_pt, text_pt, f, typeface, PP_ALIGN.LEFT)


# Symbols


def symbols(slide, field: Field, symbols_: list[Symbol], f: Palette, typeface: str) -> None:
    """Icons in light circles with their term below, in the rows that keep the terms largest, each row centred."""
    typeface = typeface or "Calibri"
    g = grid.base_pt(field.size_pt)
    left, top, width, height = _metrics(field)
    gap = 1.5 * grid.GAP * g

    def measure(rows: list[int]):
        column = width / max(rows)
        inner = column - g

        def at(k: float):
            text_pt = g * k
            d = min(_SYMBOL_CIRCLE * g * k, column * _SYMBOL_CIRCLE_SHARE)
            lines = max(_lines(s.label, typeface, text_pt, inner, True) for s in symbols_)
            row_ = d + 0.5 * text_pt + lines * text_pt * LINE_RATIO
            h = len(rows) * row_ + (len(rows) - 1) * gap
            whole = _whole([s.label for s in symbols_], font_for(typeface, text_pt, True), inner)
            return whole and h <= height, (rows, column, inner, text_pt, d, row_, h)

        return at

    # The widest grid comes first, one with more rows wins only with larger terms.
    best = None
    columns_max = max(1, int(width // (grid.SYMBOL_MIN_WIDTH * g)))
    for columns in range(min(len(symbols_), columns_max), 0, -1):
        attempt = _largest(measure(grid.distribution(len(symbols_), columns)), g, 1.0)
        if _better(attempt, best):
            best = attempt
    rows, column, inner, text_pt, d, row_, h = best[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    color_hex = str(f.tint(fb.FULL))
    icon = d * _ICON_IN_CIRCLE
    k = 0
    for r, count in enumerate(rows):
        start = left + (width - count * column) / 2
        y = top + r * (row_ + gap)
        for j, s in enumerate(symbols_[k:k + count]):
            middle = start + column * (j + 0.5)
            _shape(slide, MSO_SHAPE.OVAL, middle - d / 2, y, d, d, f.tint(fb.AREA))
            slide.shapes.add_picture(icons.picture(s.icon, color_hex), _emu(middle - icon / 2), _emu(y + (d - icon) / 2),
                                     _emu(icon), _emu(icon))
            bottom = y + d + 0.5 * text_pt
            _text(slide, middle - inner / 2, bottom, inner, y + row_ - bottom, s.label, text_pt, _BLACK, True,
                  typeface)
        k += count
