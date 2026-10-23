"""Tables and charts in the template's colour, sized to whatever count of rows, columns and series they hold."""

from __future__ import annotations

import logging

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import colors as fb
from .colors import Palette
from .schemas import MAX_TEXT_ITEMS, Body
from .template import EMU_PER_PT, Field
from .fonts import load_font

log = logging.getLogger(__name__)

_CHART = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}
# PowerPoint's own "No Style, No Grid", so the fills set here are the only ones drawn.
_NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
LINE_RATIO = 1.2
# A line that fills its box exactly wraps in PowerPoint, so text is laid out on this share of its box.
_SLACK = 0.95
# Below this even a table of eight columns is no longer readable from the back of a room.
MIN_PT = 10.0
# A template with large body text would blow a small table up to poster size.
MAX_PT = 20.0
# Chart text a notch below the body text, never tiny and never louder than the heading.
_CHART_RATIO = 0.85
_CHART_PT = (10.0, 18.0)
# Value labels crowd a chart beyond this many points, the axis then reads better.
_MAX_LABELLED_POINTS = 12
# A cell's air in multiples of the text size: at least the snug one, the roomy one where the table has space.
_SIDE = (0.5, 0.9)
_TOP = (0.3, 0.75)


def hard_lines(text: str, font, width_pt: float) -> int:
    """Greedy word wrap that also breaks a word too long for one line, as PowerPoint does."""
    if width_pt <= 0:
        return 1
    n, current = 1, 0.0
    blank = font.getlength(" ")
    for word in text.split():
        w = font.getlength(word)
        if current and current + blank + w > width_pt:
            n += 1
            current = 0.0
        if w > width_pt:
            n += int(w // width_pt)
            w = w % width_pt
        current += (blank if current else 0.0) + w
    return n


def font_for(typeface: str, size: float, bold: bool = False):
    name = f"{typeface}:bold" if bold else typeface
    return load_font(name, round(size))


def style_run(run, size: float, color: RGBColor | None = None, bold: bool = False, typeface: str = "") -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if typeface:
        run.font.name = typeface


def table(slide, field: Field, body: Body, f: Palette, typeface: str) -> None:
    """Header in full colour, rows banded at five percent, widths and size fitted to the content."""
    head = body.header[:MAX_TEXT_ITEMS]
    columns = min(MAX_TEXT_ITEMS, max([len(head)] + [len(r) for r in body.rows]))
    cells = ([head + [""] * (columns - len(head))] if head else []) + [
        (r + [""] * columns)[:columns] for r in body.rows]
    size, widths, heights, side, top = _table_metrics(cells, field, typeface or "Calibri", bool(head))
    frame = slide.shapes.add_table(len(cells), columns, field.left, field.top, field.width, sum(heights))
    t = frame.table
    t._tbl.tblPr.find(qn("a:tableStyleId")).text = _NO_STYLE
    t.first_row = bool(head)
    t.horz_banding = False
    for c, b in enumerate(widths):
        t.columns[c].width = Emu(b)
    for r, h in enumerate(heights):
        t.rows[r].height = Emu(h)
    for r, row in enumerate(cells):
        is_header = bool(head) and r == 0
        band = r - (1 if head else 0)
        for c, text in enumerate(row):
            cell = t.cell(r, c)
            cell.fill.solid()
            if is_header:
                cell.fill.fore_color.rgb = f.tint(fb.FULL)
            elif band % 2 == 1:
                cell.fill.fore_color.rgb = f.tint(fb.BAND)
            else:
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Emu(int(size * side * EMU_PER_PT))
            cell.margin_top = cell.margin_bottom = Emu(int(size * top * EMU_PER_PT))
            cell.text = text
            for run in cell.text_frame.paragraphs[0].runs:
                style_run(run, size, f.text_on() if is_header else RGBColor(0, 0, 0), is_header, typeface)


def _column_widths(cells: list[list[str]], total_pt: float, font, margin_pt: float) -> list[float] | None:
    """Widths by text length when all fits one line, else each column its longest word first, None when even that overflows."""
    n = len(cells[0])
    word = [max((font.getlength(w) for z in cells for w in z[c].split()), default=0.0) / _SLACK + margin_pt
            for c in range(n)]
    rest = total_pt - sum(word)
    if rest < 0:
        return None
    one_line = [max(font.getlength(z[c]) for z in cells) / _SLACK + margin_pt for c in range(n)]
    if sum(one_line) <= total_pt:
        # Every cell fits one line, so the columns simply share the width by their length.
        return [total_pt * l / sum(one_line) for l in one_line]
    # Long text may wrap, so a column claims at most twice the mean however long its sentence.
    cap = 2 * sum(one_line) / n
    need = [max(0.0, min(l, cap) - w) + 1 for l, w in zip(one_line, word)]
    return [w + rest * b / sum(need) for w, b in zip(word, need)]


def _table_metrics(cells: list[list[str]], field: Field, typeface: str,
                  head: bool) -> tuple[float, list[int], list[int], float, float]:
    """The largest size at which every row fits the field and no word breaks, with widths, heights and cell air.

    The roomy side air is taken when it breaks no extra line, the air above and below grows with the
    room the rows leave free.
    """
    width_pt, height_pt = field.width / EMU_PER_PT, field.height / EMU_PER_PT
    size = max(MIN_PT, float(min(MAX_PT, round(field.size_pt or 16))))
    while True:
        # Measured in bold, so the header row fits as well as the body.
        font = font_for(typeface, size, bold=head)
        tight = _column_widths(cells, width_pt, font, 2 * _SIDE[0] * size)
        if tight is None and size > MIN_PT:
            size -= 1
            continue
        tight = tight or [width_pt / len(cells[0])] * len(cells[0])
        lines = _line_counts(cells, tight, size, _SIDE[0], typeface, head)
        widths, side = tight, _SIDE[0]
        wide = _column_widths(cells, width_pt, font, 2 * _SIDE[1] * size)
        if wide is not None and _line_counts(cells, wide, size, _SIDE[1], typeface, head) == lines:
            widths, side = wide, _SIDE[1]
        text = sum(lines) * size * LINE_RATIO
        if text + 2 * _TOP[0] * size * len(cells) <= height_pt or size <= MIN_PT:
            free = (height_pt - text) / (2 * size * len(cells))
            top = max(_TOP[0], min(_TOP[1], free))
            if free < _TOP[0]:
                log.info("The table runs over its field even at %s pt", MIN_PT)
            emu = [int(b * EMU_PER_PT) for b in widths]
            emu[-1] += field.width - sum(emu)
            heights = [int((n * LINE_RATIO + 2 * top) * size * EMU_PER_PT) for n in lines]
            return size, emu, heights, side, top
        size -= 1


def _line_counts(cells: list[list[str]], widths: list[float], size: float, side: float, typeface: str,
                  head: bool) -> list[int]:
    """Lines per row, the header measured in bold."""
    out = []
    for r, row in enumerate(cells):
        font = font_for(typeface, size, bold=head and r == 0)
        out.append(max(hard_lines(t, font, (b - 2 * side * size) * _SLACK) for t, b in zip(row, widths)))
    return out


def chart(slide, field: Field, body: Body, f: Palette, typeface: str) -> None:
    """Series and slices in steps of the one colour, light grid, a legend only where it tells something."""
    data = CategoryChartData()
    data.categories = body.categories
    for series_entry in body.series:
        data.add_series(series_entry.name or " ", series_entry.values)
    kind = body.chart_type if body.chart_type in _CHART else "bar"
    chart = slide.shapes.add_chart(_CHART[kind], field.left, field.top, field.width, field.height, data).chart
    size = min(_CHART_PT[1], max(_CHART_PT[0], round((field.size_pt or 16) * _CHART_RATIO)))
    chart.font.size = Pt(size)
    chart.font.color.rgb = RGBColor.from_string(fb.SECONDARY_TEXT)
    if typeface:
        chart.font.name = typeface
    chart.has_title = False
    plot = chart.plots[0]
    points = len(body.categories) * len(body.series)
    if kind == "pie":
        _pie(chart, plot, len(body.categories), f)
    else:
        _series(chart, plot, kind, len(body.series), f)
        if points <= _MAX_LABELLED_POINTS:
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format = "General"
            labels.number_format_is_linked = False
            labels.position = XL_LABEL_POSITION.OUTSIDE_END if kind == "bar" else XL_LABEL_POSITION.ABOVE
            labels.font.color.rgb = RGBColor(0, 0, 0)
            labels.font.size = Pt(size)


def _series(chart, plot, kind: str, n: int, f: Palette) -> None:
    for series_entry, color in zip(plot.series, f.levels(n)):
        if kind == "bar":
            series_entry.format.fill.solid()
            series_entry.format.fill.fore_color.rgb = color
            series_entry.format.line.fill.background()
        else:
            series_entry.smooth = False
            series_entry.format.line.color.rgb = color
            series_entry.format.line.width = Pt(2.5)
            series_entry.marker.style = XL_MARKER_STYLE.CIRCLE
            series_entry.marker.size = 7
            series_entry.marker.format.fill.solid()
            series_entry.marker.format.fill.fore_color.rgb = color
            series_entry.marker.format.line.color.rgb = color
    if kind == "bar":
        plot.gap_width = 60
        plot.overlap = 0
    chart.has_legend = n > 1
    if n > 1:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    grey = RGBColor.from_string(fb.GRIDLINE)
    value = chart.value_axis
    value.has_major_gridlines = True
    value.major_gridlines.format.line.color.rgb = grey
    value.major_gridlines.format.line.width = Pt(0.75)
    value.format.line.fill.background()
    chart.category_axis.format.line.color.rgb = grey


def _pie(chart, plot, n: int, f: Palette) -> None:
    """Slices from full colour down to the lightest step, split by white edges, named in the legend."""
    for point, color in zip(plot.series[0].points, f.levels(n)):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = color
        point.format.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        point.format.line.width = Pt(1.5)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = "General"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.color.rgb = RGBColor(0, 0, 0)
