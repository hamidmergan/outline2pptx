"""The library: an outline from a topic, and a deck from an outline."""

from __future__ import annotations

import logging
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Callable, Optional

from . import colors, library, llm, outline as ol, pictures, slides as sl, split, template
from .build import build_deck, date_text
from .config import Settings, resolve
from .details import details
from .errors import ModelError
from .outline import Unusable
from .schemas import Outline, OutlineRow, DeckParams, Slide

log = logging.getLogger(__name__)

# How often one call is repeated when the answer comes back unusable.
OUTLINE_TRIES = 2
SLIDE_TRIES = 3

Progress = Optional[Callable[[str, int, int], None]]


class Deck:
    """A finished deck: its file, its slides and what it was made with."""

    def __init__(self, data: bytes, outline: Outline, slides: list[Slide], info: dict):
        self.data = data
        self.outline = outline
        self.slides = slides
        self.info = info

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(self.data)
        return path

    def preview(self, dpi: int | None = None) -> list[bytes]:
        """One PNG per slide. Needs LibreOffice on the machine."""
        from . import preview as pv

        return pv.render(self.data, dpi or pv.DPI)


def _params(settings: Settings, meta: dict, fallback: bool = False) -> DeckParams:
    return DeckParams(language=settings.language, instructions=settings.instructions,
                      presenter=settings.presenter, notes=settings.notes,
                      template_id=meta.get("id", settings.template),
                      template_name=meta.get("name", ""), template_fallback=fallback,
                      style_text=settings.style_text or meta.get("style_text", ""))


def list_templates(template_dir: str | Path | None = None) -> list[dict]:
    """Every template that can be used, by id, with its name and picture style."""
    return [{k: v for k, v in t.items() if k != "path"} for t in library.available(template_dir)]


def outline_from_topic(topic: str, *, slides: int | None = None, client=None, **options) -> str:
    """An outline for a topic, as Markdown, planned by the model."""
    settings = resolve(slides=slides, **options)
    if not topic.strip():
        raise ModelError("No topic was given.")
    _, analysis, meta = library.load(settings.template, settings.template_dir)
    params = _params(settings, meta)
    roles = template.selectable_roles(analysis)
    client = client or llm.make_client(settings)
    last = ""
    for _ in range(OUTLINE_TRIES):
        messages = ol.build_messages(topic.strip(), params, roles, settings.slides)
        answer = llm.call_json(client, settings, messages)
        try:
            return ol.to_markdown(ol.parse_answer(answer, roles))
        except Unusable as exc:
            last = str(exc)
            log.warning("The outline was unusable, asking again: %s", exc)
    raise ModelError(last or "The model returned no usable outline.")


def build(outline_md: str, output: str | Path | None = None, *, client=None,
          progress: Progress = None, **options) -> Deck:
    """A deck from a Markdown outline: the model writes every slide, python-pptx builds the file."""
    settings = resolve(**options)
    data, analysis, meta = library.load(settings.template, settings.template_dir)
    params = _params(settings, meta)
    roles = template.selectable_roles(analysis)
    plan = ol.from_markdown(outline_md, roles)
    client = client or llm.make_client(settings)

    layouts = _layouts(plan, analysis)
    written = _write(plan, layouts, analysis, params, settings, client, progress)
    ordered = split.split_overflow([written[r.id] for r in plan.rows if r.id in written], analysis)
    layouts.update({s.id: s.layout for s in ordered if s.layout is not None})

    palette = colors.from_template(_presentation(data), analysis)
    image_client = None
    if settings.images and settings.has_image_model:
        image_client = llm.make_image_client(settings)
    if progress:
        progress("pictures", 0, len(ordered))
    drawn = pictures.draw(analysis, ordered, layouts, settings, params.style_text, palette.base,
                          client=image_client)

    deck = build_deck(data, analysis, ordered, drawn, plan.title, params.notes_in_deck,
                      settings.notice, date=date_text(settings.language),
                      presenter=settings.presenter)
    info = details(params, settings.model, settings.image_model if settings.images else "")
    info["slides"] = len(ordered)
    result = Deck(deck, plan, ordered, info)
    if output:
        result.save(output)
    if progress:
        progress("done", len(ordered), len(ordered))
    return result


def _presentation(data: bytes):
    from io import BytesIO

    from pptx import Presentation

    return Presentation(BytesIO(data))


def _layouts(plan: Outline, analysis) -> dict[str, int]:
    """One layout per row, rotated so neighbours of one role look different."""
    counter: dict[str, int] = {}
    return {r.id: template.layout_for(analysis, r.role, counter).index for r in plan.rows}


def _write(plan: Outline, layouts: dict[str, int], analysis, params: DeckParams,
           settings: Settings, client, progress: Progress) -> dict[str, Slide]:
    rows = plan.rows
    batches = [rows[i:i + settings.batch] for i in range(0, len(rows), settings.batch)]
    # Planned over the whole deck, the batches run at once and cannot see each other's cards.
    all_specs = {r.id: template.spec_for(analysis.layout(layouts[r.id]), analysis) for r in rows}
    with_cards = sl.card_slides(rows, all_specs)
    source = f"{plan.title}. {plan.subtitle}".strip()
    done = 0
    out: dict[str, Slide] = {}

    def write(batch: list[OutlineRow]) -> list[Slide]:
        specs = {r.id: all_specs[r.id] if r.id in with_cards else sl.without_cards(all_specs[r.id])
                 for r in batch}
        messages = sl.build_messages(source, params, batch, specs)
        last = ""
        for _ in range(SLIDE_TRIES):
            answer = llm.call_json(client, settings, messages)
            try:
                fresh = sl.parse_answer(answer, batch, specs, params.notes)
            except Unusable as exc:
                last = str(exc)
                log.warning("A batch of slides came back unusable: %s", exc)
                continue
            return _ask_again(client, settings, messages, answer, fresh, batch, specs, params.notes)
        raise ModelError(last or "The model returned no usable slides.")

    waves = [batches[i:i + settings.write_parallelism]
             for i in range(0, len(batches), settings.write_parallelism)]
    for wave in waves:
        with ThreadPoolExecutor(max_workers=len(wave)) as pool:
            results = list(pool.map(write, wave))
        for fresh in results:
            for s in fresh:
                s.layout = layouts[s.id]
                out[s.id] = s
            done += len(fresh)
        if progress:
            progress("slides", done, len(rows))
    return out


def _ask_again(client, settings: Settings, messages, answer, fresh: list[Slide],
               batch: list[OutlineRow], specs: dict, notes: bool) -> list[Slide]:
    """Thin slides are asked for once more, the second version wins only when it carries more points."""
    thin = sl.too_thin(fresh, specs)
    if not thin:
        return fresh
    log.info("Too few points, asking again: %s", thin)
    try:
        raw = llm.call_json(client, settings, sl.ask_again(messages, answer, thin))
        second = {s.id: s for s in sl.parse_answer(raw, [r for r in batch if r.id in thin],
                                                   specs, notes)}
    except Exception as exc:  # noqa: BLE001 - the first version is still a usable slide
        log.warning("The second ask failed, the first version stays: %s", exc)
        return fresh
    return [second[s.id] if s.id in second and sl.bullet_count(second[s.id]) > sl.bullet_count(s)
            else s for s in fresh]
