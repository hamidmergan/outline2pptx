"""Plain text out of the document formats a deck is built from, from bytes or from a path."""

import re
from io import BytesIO
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

Source = bytes | str | Path


def _as_bytes(src: Source) -> bytes:
    """Raw bytes for the text formats, which have to decode the whole thing anyway."""
    if isinstance(src, bytes):
        return src
    with open(src, "rb") as f:
        return f.read()


def _as_stream(src: Source):
    """What python-docx / python-pptx should open: a path stays a path."""
    return BytesIO(src) if isinstance(src, bytes) else src


def extract_text_md(content: Source) -> str:
    # Replacing, like the other text readers: a file in an older encoding is still read
    # with a few broken signs instead of being refused as unreadable.
    text = _as_bytes(content).decode("utf-8", errors="replace")
    # Remove code blocks
    text = re.sub(r"```[\s\S]*?```", "", text)
    text = re.sub(r"`[^`]+`", "", text)
    # Remove images
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    # Convert links to just the link text
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    # Remove headings markers
    text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
    # Remove bold/italic markers
    text = re.sub(r"\*{1,3}([^*]+)\*{1,3}", r"\1", text)
    text = re.sub(r"_{1,3}([^_]+)_{1,3}", r"\1", text)
    # Remove horizontal rules
    text = re.sub(r"^[-*_]{3,}\s*$", "", text, flags=re.MULTILINE)
    # Remove HTML tags
    text = re.sub(r"<[^>]+>", "", text)
    # Collapse multiple blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_docx(content: Source) -> str:
    doc = Document(_as_stream(content))
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return "\n".join(paragraphs)


# Footer, slide number and date stand on every slide and say nothing about it.
_PPTX_FURNITURE = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}


def pptx_slide_texts(content: Source) -> list[str]:
    """The text of every slide on its own, without footer, slide number and date."""
    prs = Presentation(_as_stream(content))
    slides = []
    for slide in prs.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.type in _PPTX_FURNITURE:
                continue
            _collect_shape_text(shape, texts)
        slides.append("\n".join(texts))
    return slides


def extract_text_pptx(content: Source) -> str:
    return "\n".join(t for t in pptx_slide_texts(content) if t)


def _collect_shape_text(shape, texts: list) -> None:
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:
                texts.append(text)
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for subshape in shape.shapes:
            _collect_shape_text(subshape, texts)


def extract_text_txt(content: Source) -> str:
    return _as_bytes(content).decode("utf-8", errors="replace")


EXTRACTORS = {
    ".txt": extract_text_txt,
    ".md": extract_text_md,
    ".docx": extract_text_docx,
    ".pptx": extract_text_pptx,
}
