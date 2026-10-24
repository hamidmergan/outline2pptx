"""One input file into the canonical text of a run, with its pages.

What comes out of here is the single truth about a source for the rest of the run. The
text is produced once and never derived again, because the deck is planned over exactly
this text.
"""

from __future__ import annotations

import logging
import threading
from dataclasses import dataclass
from pathlib import Path

from . import config
from .errors import LimitError, NoReadableText, UnsupportedFileError
from .pdf_layout import (
    boilerplate_keys,
    line_boxes,
    merge_rows,
    page_reading_order,
    strip_boilerplate_text,
    strip_furniture,
)
from .text_clean import clean

log = logging.getLogger(__name__)

# One pdfium call at a time per process, the library is not thread safe.
PDFIUM_LOCK = threading.RLock()

# What a page break looks like in the assembled text.
_PAGE_BREAK = "\n\n"

# The formats a deck can be built from.
EXTENSIONS: tuple[str, ...] = (".pdf", ".docx", ".pptx", ".md", ".txt")


@dataclass(frozen=True)
class Source:
    """A read file as the rest of the run sees it."""

    name: str
    text: str
    pages: int


def _pdf_pages(path: str) -> list[str]:
    """Every page of a PDF as cleaned text, headers and footers counted out.

    No text recognition: a page without a text layer stays empty and the caller
    refuses the file rather than guessing at a picture.
    """
    import pypdfium2 as pdfium

    page_texts: list[str] = []
    page_boxes: list[list] = []
    with PDFIUM_LOCK:
        doc = pdfium.PdfDocument(path)
        try:
            for number in range(len(doc)):
                page = doc[number]
                textpage = page.get_textpage()
                page_texts.append(clean(textpage.get_text_range() or "").strip())
                page_boxes.append(line_boxes(textpage))
                textpage.close()
                page.close()
        finally:
            doc.close()

    # Counted over the whole document, a header only shows itself in the whole.
    boiler = boilerplate_keys(page_boxes, page_texts)
    out: list[str] = []
    for boxes, raw in zip(page_boxes, page_texts):
        ordered = page_reading_order(strip_furniture(boxes, boiler))
        rows = merge_rows(ordered)
        out.append("\n".join(b.text for b in rows) if rows
                   else strip_boilerplate_text(raw, boiler))
    log.info("Source %s: %d page(s), %d header or footer line(s) removed",
             Path(path).name, len(out), len(boiler))
    return out


def _pptx_page(path: str) -> str:
    """All slides as one page, the lines that stand on most slides counted out."""
    from .extract import pptx_slide_texts

    slides = [clean(s).strip() for s in pptx_slide_texts(path)]
    boiler = boilerplate_keys([], slides)
    log.info("Source %s: %d slide(s), %d header or footer line(s) removed",
             Path(path).name, len(slides), len(boiler))
    return "\n".join(strip_boilerplate_text(s, boiler) for s in slides if s)


def _other_pages(path: str, ext: str) -> list[str]:
    """Every other format as one single page, they carry no page numbers worth keeping."""
    from .extract import EXTRACTORS

    if ext == ".pptx":
        return [_pptx_page(path)]
    reader = EXTRACTORS.get(ext)
    if reader is None:
        raise UnsupportedFileError(f"Format {ext} is not supported")
    return [clean(reader(path) or "").strip()]


def _assemble(pages: list[str]) -> str:
    """The pages as one text. An empty page takes no room."""
    return _PAGE_BREAK.join(page for page in (raw.strip() for raw in pages) if page)


def count_pages(ext: str, path: str) -> int | None:
    """Page or slide count of a file on disk, None where the format has no pages."""
    try:
        if ext == ".pdf":
            import pypdfium2 as pdfium
            with PDFIUM_LOCK:
                doc = pdfium.PdfDocument(path)
                try:
                    return len(doc)
                finally:
                    doc.close()
        if ext == ".pptx":
            from pptx import Presentation
            return len(Presentation(path).slides)
    except Exception as e:  # noqa: BLE001 - an unreadable count must not fail the file
        log.warning("Page count of %s not readable: %s", ext, e)
    return None


def read_file(path: str | Path, name: str | None = None, *,
              max_pages: int = config.MAX_SOURCE_PAGES,
              max_file_mb: float = config.MAX_FILE_MB) -> Source:
    """One file to the canonical source record. Raises when it carries no text."""
    path = str(path)
    name = name or Path(path).name
    ext = Path(name).suffix.lower()
    if ext not in EXTENSIONS:
        raise UnsupportedFileError(
            f"{name}: format {ext or '(none)'} is not supported. Supported: {', '.join(EXTENSIONS)}")
    p = Path(path)
    if not p.is_file():
        raise UnsupportedFileError(f"{name}: file not found")
    size_mb = p.stat().st_size / (1024 * 1024)
    if size_mb > max_file_mb:
        raise LimitError(f"{name}: the file has {size_mb:.1f} MB, at most {max_file_mb:g} MB are "
                         "allowed (max_file_mb or OUTLINE2PPTX_MAX_FILE_MB).")
    if p.stat().st_size == 0:
        raise UnsupportedFileError(f"{name}: the file is empty")
    pages = count_pages(ext, path)
    if pages is not None and pages > max_pages:
        raise LimitError(f"{name}: the document has {pages} pages, at most {max_pages} are allowed "
                         "(max_pages or OUTLINE2PPTX_MAX_SOURCE_PAGES).")
    try:
        page_texts = _pdf_pages(path) if ext == ".pdf" else _other_pages(path, ext)
    except UnsupportedFileError:
        raise
    except Exception as e:  # noqa: BLE001 - a broken file is no fault of the run
        raise UnsupportedFileError(f"{name}: the file cannot be read ({e})") from e
    text = _assemble(page_texts)
    if not text.strip():
        raise NoReadableText(f"{name}: the file carries no readable text, a scan without a text "
                             "layer for instance.")
    log.info("Source %s: %d characters", name, len(text))
    return Source(name=name, text=text, pages=len(page_texts))


def from_text(name: str, text: str) -> Source:
    """A text with no document behind it, it becomes the same record a file becomes."""
    whole = _assemble([clean(text or "")])
    if not whole.strip():
        raise NoReadableText(f"{name}: the text is empty")
    return Source(name=name, text=whole, pages=1)
