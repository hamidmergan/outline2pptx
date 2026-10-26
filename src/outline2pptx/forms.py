"""Timeline, cycle, comparison, matrix, quote, checklist, definitions, mistakes, scale, Venn
diagram, tree and worked example, each drawn from shapes and sized to the text it holds."""

from __future__ import annotations

import math

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from . import colors as fb
from . import grid, icons
from .colors import Palette
from .diagrams import _TOP_SHARE, _entry, _entry_height, _largest, _lines, _polygon, _text, _whole
from .elements import LINE_RATIO, MIN_PT, font_for, style_run
from .schemas import Body, Branch, Cell, Entry, Event, Mark, Pair, Side
from .shapes import _BLACK, _CORNER, _emu, _shape, _parts, _textbox
from .template import EMU_PER_PT, Field

# The colours of wrong and right, the same in every template so they read the same everywhere.
BAD = "B91C1C"
GOOD = "15803D"
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# Air inside a panel and between rows, in multiples of the base size.
_PAD = 0.8
_AIR = 0.5
_LINE_PT = 1.5
# A quotation sits in the size of a statement, never louder than this.
_QUOTE_PT = 36.0
_QUOTE_RATIO = 1.9
# The ring of a cycle takes at most this share of the field's width, a Venn circle this share of the height.
_RING_SHARE = 0.42
_VENN_HEIGHT = 0.95
_VENN_OVERLAP = 0.34


def _metrics(field: Field) -> tuple[float, float, float, float]:
    return field.left / EMU_PER_PT, field.top / EMU_PER_PT, field.width / EMU_PER_PT, field.height / EMU_PER_PT


def _h(text: str, size: float, width: float, typeface: str, bold: bool = False) -> float:
    return _lines(text, typeface, size, width, bold) * size * LINE_RATIO if text else 0.0


def _sizes(g: float, k: float) -> tuple[float, float]:
    return grid.TITLE_RATIO * g * k, grid.TEXT_RATIO * g * k


def _fits(texts: list[str], typeface: str, size: float, width: float, bold: bool = False) -> bool:
    return _whole([t for t in texts if t], font_for(typeface, size, bold), width)


def _rgb(hex_value: str, ratio: float = fb.FULL) -> RGBColor:
    return RGBColor.from_string(fb.mix(hex_value, ratio))


def _panel(slide, x: float, y: float, w: float, h: float, color, round_: bool = True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    return _shape(slide, kind, x, y, w, h, color, _CORNER if round_ else None)


def _number(slide, x: float, y: float, d: float, n: int, f: Palette, typeface: str) -> None:
    """A filled circle in the colour with its number."""
    _shape(slide, MSO_SHAPE.OVAL, x, y, d, d, f.tint(fb.FULL))
    tf = _text(slide, x, y, d, d, str(n), d * 0.45, f.text_on(), True, typeface, anchor=MSO_ANCHOR.MIDDLE)
    tf.word_wrap = False


def _icon(slide, name: str, color: str, x: float, y: float, size: float) -> None:
    slide.shapes.add_picture(icons.picture(name, color), _emu(x), _emu(y), _emu(size), _emu(size))


def _rich(slide, x: float, y: float, w: float, h: float, text: str, marked: str, size: float, f: Palette,
          typeface: str, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT) -> None:
    """A sentence with its one term in the colour and bold."""
    tf = _textbox(slide, x, y, w, h, anchor, align)
    for part, on in _parts(text, marked):
        run = tf.paragraphs[0].add_run()
        run.text = part
        style_run(run, size, f.text_color() if on else _BLACK, on, typeface)


def _start(field: Field, typeface: str) -> tuple[str, float, float, float, float, float]:
    left, top, width, height = _metrics(field)
    return typeface or "Calibri", grid.base_pt(field.size_pt), left, top, width, height


# Timeline


def timeline(slide, field: Field, events: list[Event], f: Palette, typeface: str) -> None:
    """Dates above an arrow in reading order, a dot on it for each event, its title and sentence below."""
    typeface, g, left, top, width, height = _start(field, typeface)
    n = len(events)
    column = width / n
    inner = column - g

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        date_pt = title_pt * 1.15
        d = 0.9 * g * k
        date_h = max(_h(e.date, date_pt, inner, typeface, True) for e in events)
        block = max(_entry_height(e, typeface, title_pt, text_pt, inner) for e in events)
        h = date_h + 0.5 * g * k + d + 0.6 * g * k + block
        whole = (_fits([e.date for e in events], typeface, date_pt, inner, True)
                 and _fits([e.title for e in events], typeface, title_pt, inner, True)
                 and _fits([e.text for e in events], typeface, text_pt, inner))
        return whole and h <= height, (title_pt, text_pt, date_pt, d, date_h, h, k)

    title_pt, text_pt, date_pt, d, date_h, h, k = _largest(measure, g)[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    line_y = top + date_h + 0.5 * g * k
    thickness = max(2.0, d * 0.22)
    head = d * 0.9
    _shape(slide, MSO_SHAPE.RECTANGLE, left, line_y + (d - thickness) / 2, width - head, thickness, f.tint(fb.LINE))
    _polygon(slide, [(left + width - head, line_y + d / 2 - head / 2), (left + width, line_y + d / 2),
                     (left + width - head, line_y + d / 2 + head / 2)], f.tint(fb.LINE))
    for i, e in enumerate(events):
        cx = left + column * (i + 0.5)
        tf = _text(slide, cx - inner / 2, top, inner, date_h, e.date, date_pt, f.text_color(), True, typeface,
                   anchor=MSO_ANCHOR.BOTTOM)
        tf.word_wrap = True
        _shape(slide, MSO_SHAPE.OVAL, cx - d / 2, line_y, d, d, f.tint(fb.FULL))
        _entry(slide, cx - inner / 2, line_y + d + 0.6 * g * k, inner, e, title_pt, text_pt, f, typeface)


# Cycle


def cycle(slide, field: Field, stages: list[Entry], center: str, f: Palette, typeface: str) -> None:
    """A ring with a numbered dot per stage and the cycle's name inside, the stages listed beside it."""
    typeface, g, left, top, width, height = _start(field, typeface)
    n = len(stages)
    # The numbered dots stick out over the ring, so the ring keeps clear of the field's edges.
    ring = min(height * 0.85, width * _RING_SHARE)
    text_x = left + ring + 2 * g
    text_w = left + width - text_x

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        badge = 1.3 * title_pt
        inner = text_w - badge - 0.6 * g
        blocks = [max(badge, _entry_height(e, typeface, title_pt, text_pt, inner)) for e in stages]
        air = _AIR * g * k
        h = sum(blocks) + air * (n - 1)
        whole = (_fits([e.title for e in stages], typeface, title_pt, inner, True)
                 and _fits([e.text for e in stages], typeface, text_pt, inner))
        return whole and h <= height, (title_pt, text_pt, badge, inner, blocks, air, h)

    title_pt, text_pt, badge, inner, blocks, air, h = _largest(measure, g)[2]
    ring_y = top + max(0.0, (height - ring) * _TOP_SHARE)
    donut = _shape(slide, MSO_SHAPE.DONUT, left, ring_y, ring, ring, f.tint(fb.LINE))
    donut.adjustments[0] = 0.08
    node = min(ring * 0.2, 2.6 * g)
    radius = ring / 2 - ring * 0.04
    cx, cy = left + ring / 2, ring_y + ring / 2
    for i in range(n):
        angle = -math.pi / 2 + 2 * math.pi * i / n
        _number(slide, cx + radius * math.cos(angle) - node / 2, cy + radius * math.sin(angle) - node / 2, node,
                i + 1, f, typeface)
    if center:
        inside = ring * 0.5
        size = title_pt * 1.2
        while size > MIN_PT and (_h(center, size, inside, typeface, True) > inside
                                 or not _fits([center], typeface, size, inside, True)):
            size -= 1
        _text(slide, cx - inside / 2, cy - inside / 2, inside, inside, center, size, f.text_color(), True,
              typeface, anchor=MSO_ANCHOR.MIDDLE)
    y = top + max(0.0, (height - h) * _TOP_SHARE)
    for i, (e, block) in enumerate(zip(stages, blocks)):
        _number(slide, text_x, y, badge, i + 1, f, typeface)
        own = _entry_height(e, typeface, title_pt, text_pt, inner)
        _entry(slide, text_x + badge + 0.6 * g, y + max(0.0, (badge - own) / 2), inner, e, title_pt, text_pt, f,
               typeface, PP_ALIGN.LEFT)
        y += block + air


# Versus


def versus(slide, field: Field, sides: list[Side], tone: str, f: Palette, typeface: str) -> None:
    """Two panels side by side, each with its name in a coloured head and its points below, a badge between them."""
    typeface, g, left, top, width, height = _start(field, typeface)
    gap = 3 * g
    panel_w = (width - gap) / 2
    if tone == "good_bad":
        heads = [_rgb(BAD), _rgb(GOOD)]
        marks = [BAD, GOOD]
    else:
        heads = [f.tint(fb.FULL), f.tint(0.65)]
        marks = [str(f.tint(fb.FULL))] * 2

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        pad = _PAD * g * k
        dot = 0.4 * text_pt
        inner = panel_w - 2 * pad - dot - 0.5 * text_pt
        head_h = max(_h(s.title, title_pt, panel_w - 2 * pad, typeface, True) for s in sides) + 1.2 * pad
        items_h = max(sum(_h(t, text_pt, inner, typeface) + 0.5 * text_pt for t in s.items) for s in sides)
        h = head_h + pad + items_h + pad
        whole = (_fits([s.title for s in sides], typeface, title_pt, panel_w - 2 * pad, True)
                 and _fits([t for s in sides for t in s.items], typeface, text_pt, inner))
        return whole and h <= height, (title_pt, text_pt, pad, dot, inner, head_h, h)

    title_pt, text_pt, pad, dot, inner, head_h, h = _largest(measure, g)[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    for i, (side, head) in enumerate(zip(sides, heads)):
        x = left + i * (panel_w + gap)
        _panel(slide, x, top, panel_w, h, f.tint(fb.AREA))
        _panel(slide, x, top, panel_w, head_h, head)
        on = _WHITE if tone == "good_bad" else (f.text_on() if i == 0 else f.text_on(0.65))
        _text(slide, x + pad, top, panel_w - 2 * pad, head_h, side.title, title_pt, on, True, typeface,
              anchor=MSO_ANCHOR.MIDDLE)
        y = top + head_h + pad
        for item in side.items:
            item_h = _h(item, text_pt, inner, typeface)
            _shape(slide, MSO_SHAPE.OVAL, x + pad, y + (text_pt * LINE_RATIO - dot) / 2, dot, dot, _rgb(marks[i]))
            _text(slide, x + pad + dot + 0.5 * text_pt, y, inner, item_h, item, text_pt, _BLACK, False, typeface,
                  align=PP_ALIGN.LEFT)
            y += item_h + 0.5 * text_pt
    badge = min(gap * 1.4, 2.8 * g)
    bx, by = left + panel_w + gap / 2 - badge / 2, top + h / 2 - badge / 2
    _shape(slide, MSO_SHAPE.OVAL, bx, by, badge, badge, _WHITE)
    _shape(slide, MSO_SHAPE.OVAL, bx + badge * 0.1, by + badge * 0.1, badge * 0.8, badge * 0.8, f.tint(fb.LINE))
    tf = _text(slide, bx, by, badge, badge, "vs", badge * 0.32, _BLACK, True, typeface, anchor=MSO_ANCHOR.MIDDLE)
    tf.word_wrap = False


# Matrix


def _rotated_label(slide, cx: float, cy: float, length: float, thickness: float, text: str, size: float,
                   color, typeface: str) -> None:
    """A label that reads from bottom to top, centred on the given point."""
    box = slide.shapes.add_textbox(_emu(cx - length / 2), _emu(cy - thickness / 2), _emu(length), _emu(thickness))
    box.rotation = 270
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = tf.paragraphs[0].add_run()
    run.text = text
    style_run(run, size, color, True, typeface)


def matrix(slide, field: Field, cells: list[Cell], axis_x: str, axis_y: str, f: Palette, typeface: str) -> None:
    """Four quadrants along two axes, the axes as arrows with their names, the top right one stronger."""
    typeface, g, left, top, width, height = _start(field, typeface)
    band = 2.4 * g
    gap = 0.6 * g
    gx, gw = left + band, width - band
    gh = min(height - band, gw * 0.62)
    cell_w, cell_h = (gw - gap) / 2, (gh - gap) / 2
    by_pos = {c.pos: c for c in cells}

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        pad = _PAD * g * k
        inner = cell_w - 2 * pad
        h = max(_entry_height(c, typeface, title_pt, text_pt, inner) for c in cells) + 2 * pad
        whole = (_fits([c.title for c in cells], typeface, title_pt, inner, True)
                 and _fits([c.text for c in cells], typeface, text_pt, inner))
        return whole and h <= cell_h, (title_pt, text_pt, pad, inner)

    title_pt, text_pt, pad, inner = _largest(measure, g)[2]
    gy = top + max(0.0, (height - band - gh) * _TOP_SHARE)
    spots = {"tl": (0, 0), "tr": (1, 0), "bl": (0, 1), "br": (1, 1)}
    for pos, (i, j) in spots.items():
        x, y = gx + i * (cell_w + gap), gy + j * (cell_h + gap)
        _panel(slide, x, y, cell_w, cell_h, f.tint(fb.LINE if pos == "tr" else fb.AREA))
        c = by_pos.get(pos)
        if c is not None:
            _entry(slide, x + pad, y + pad, inner, c, title_pt, text_pt, f, typeface, PP_ALIGN.LEFT)
    axis = max(2.0, 0.12 * g)
    head = 0.7 * g
    ax = gx - 0.7 * g
    _shape(slide, MSO_SHAPE.RECTANGLE, ax - axis / 2, gy + head, axis, gh - head, f.tint(fb.FULL))
    _polygon(slide, [(ax - head / 2, gy + head), (ax, gy), (ax + head / 2, gy + head)], f.tint(fb.FULL))
    ay = gy + gh + 0.7 * g
    _shape(slide, MSO_SHAPE.RECTANGLE, gx, ay - axis / 2, gw - head, axis, f.tint(fb.FULL))
    _polygon(slide, [(gx + gw - head, ay - head / 2), (gx + gw, ay), (gx + gw - head, ay + head / 2)],
             f.tint(fb.FULL))
    label_pt = text_pt
    _text(slide, gx, ay + 0.3 * g, gw, label_pt * LINE_RATIO * 1.2, axis_x, label_pt, f.text_color(), True, typeface)
    _rotated_label(slide, left + (band - 0.7 * g) / 2 - 0.2 * g, gy + gh / 2, gh, label_pt * LINE_RATIO * 1.2,
                   axis_y, label_pt, f.text_color(), typeface)


# Quote


def quote(slide, field: Field, body: Body, f: Palette, typeface: str) -> None:
    """A large quotation mark, a bar in the colour, the sentence large with its one term marked."""
    typeface, g, left, top, width, height = _start(field, typeface)
    # The glyph fills only the upper quarter of its line, so it is set large and pulled down onto the text.
    mark_pt = 7 * g
    text_x = left + 1.2 * g + 0.4 * g
    text_w = width - (text_x - left) - 2 * g
    size = min(_QUOTE_PT, _QUOTE_RATIO * g)
    while size > MIN_PT and (_h(body.text, size, text_w, typeface, True) > height * 0.75
                             or not _fits([body.text], typeface, size, text_w, True)):
        size -= 1
    block = _h(body.text, size, text_w, typeface, True)
    y = top + max(0.0, (height - block - mark_pt * 0.45) * _TOP_SHARE) + mark_pt * 0.45
    tf = _text(slide, left - 0.1 * g, y - mark_pt * 0.58, 2 * mark_pt, mark_pt, "“", mark_pt, f.tint(fb.LINE),
               True, typeface, align=PP_ALIGN.LEFT)
    tf.word_wrap = False
    _shape(slide, MSO_SHAPE.RECTANGLE, left, y, 0.35 * g, block, f.tint(fb.FULL))
    tf = _textbox(slide, text_x, y, text_w, block, MSO_ANCHOR.MIDDLE)
    for part, on in _parts(body.text, body.emphasis):
        run = tf.paragraphs[0].add_run()
        run.text = part
        style_run(run, size, f.text_color() if on else _BLACK, True, typeface)


# Checklist


def checklist(slide, field: Field, items: list[str], f: Palette, typeface: str) -> None:
    """Every item in a light box with a check mark in the colour, in one or two columns."""
    typeface, g, left, top, width, height = _start(field, typeface)
    n = len(items)
    columns = 2 if n > 4 and width >= 2 * grid.CHECK_MIN_WIDTH * g else 1
    rows = math.ceil(n / columns)
    gap = 1.5 * g
    col_w = (width - (columns - 1) * gap) / columns

    def measure(k: float):
        text_pt = g * k
        pad = 0.6 * g * k
        icon = 1.3 * text_pt
        inner = col_w - 2 * pad - icon - 0.6 * text_pt
        row_h = max(max(icon, _h(t, text_pt, inner, typeface)) for t in items) + 2 * pad
        air = 0.5 * g * k
        h = rows * row_h + (rows - 1) * air
        return _fits(items, typeface, text_pt, inner) and h <= height, (text_pt, pad, icon, inner, row_h, air, h)

    text_pt, pad, icon, inner, row_h, air, h = _largest(measure, g, 1.0)[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    color = str(f.tint(fb.FULL))
    for i, item in enumerate(items):
        c, r = divmod(i, rows)
        x, y = left + c * (col_w + gap), top + r * (row_h + air)
        _panel(slide, x, y, col_w, row_h, f.tint(fb.AREA))
        _icon(slide, "check-circle", color, x + pad, y + (row_h - icon) / 2, icon)
        _text(slide, x + pad + icon + 0.6 * text_pt, y, inner, row_h, item, text_pt, _BLACK, False, typeface,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)


# Definitions


def definitions(slide, field: Field, terms: list[Entry], f: Palette, typeface: str) -> None:
    """Terms in the colour in a narrow column, what they mean beside them, thin lines between the rows."""
    typeface, g, left, top, width, height = _start(field, typeface)
    n = len(terms)

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        widest = max(font_for(typeface, title_pt, True).getlength(e.title) for e in terms) / 0.95
        term_w = min(width * 0.32, widest + g)
        text_w = width - term_w - 1.5 * g
        air = 0.9 * g * k
        rows = [max(_h(e.title, title_pt, term_w, typeface, True), _h(e.text, text_pt, text_w, typeface))
                for e in terms]
        h = sum(rows) + air * (n - 1)
        whole = (_fits([e.title for e in terms], typeface, title_pt, term_w, True)
                 and _fits([e.text for e in terms], typeface, text_pt, text_w))
        return whole and h <= height, (title_pt, text_pt, term_w, text_w, air, rows, h)

    title_pt, text_pt, term_w, text_w, air, rows, h = _largest(measure, g)[2]
    y = top + max(0.0, (height - h) * _TOP_SHARE)
    for i, (e, row_h) in enumerate(zip(terms, rows)):
        if i:
            _shape(slide, MSO_SHAPE.RECTANGLE, left, y - air / 2 - _LINE_PT / 2, width, _LINE_PT, f.tint(fb.LINE))
        _text(slide, left, y, term_w, row_h, e.title, title_pt, f.text_color(), True, typeface, align=PP_ALIGN.LEFT)
        _text(slide, left + term_w + 1.5 * g, y, text_w, row_h, e.text, text_pt, _BLACK, False, typeface,
              align=PP_ALIGN.LEFT)
        y += row_h + air


# Mistakes


def mistakes(slide, field: Field, pairs: list[Pair], f: Palette, typeface: str) -> None:
    """Per row the mistake in a light red box and the right way in a light green one, an arrow between them."""
    typeface, g, left, top, width, height = _start(field, typeface)
    arrow = 2.4 * g
    box_w = (width - arrow) / 2

    def measure(k: float):
        text_pt = g * k
        pad = 0.7 * g * k
        icon = 1.3 * text_pt
        inner = box_w - 2 * pad - icon - 0.6 * text_pt
        rows = [max(icon, _h(z.wrong, text_pt, inner, typeface), _h(z.right, text_pt, inner, typeface)) + 2 * pad
                for z in pairs]
        air = 0.6 * g * k
        h = sum(rows) + air * (len(pairs) - 1)
        whole = _fits([t for z in pairs for t in (z.wrong, z.right)], typeface, text_pt, inner)
        return whole and h <= height, (text_pt, pad, icon, inner, rows, air, h)

    text_pt, pad, icon, inner, rows, air, h = _largest(measure, g, 1.0)[2]
    y = top + max(0.0, (height - h) * _TOP_SHARE)
    for z, row_h in zip(pairs, rows):
        for x, text, hex_value, name in ((left, z.wrong, BAD, "warning"),
                                         (left + box_w + arrow, z.right, GOOD, "check-circle")):
            _panel(slide, x, y, box_w, row_h, _rgb(hex_value, fb.AREA))
            _icon(slide, name, hex_value, x + pad, y + (row_h - icon) / 2, icon)
            _text(slide, x + pad + icon + 0.6 * text_pt, y, inner, row_h, text, text_pt, _BLACK, False, typeface,
                  align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        size = arrow * 0.55
        _shape(slide, MSO_SHAPE.RIGHT_ARROW, left + box_w + (arrow - size) / 2, y + row_h / 2 - size * 0.35, size,
               size * 0.7, f.tint(fb.LINE))
        y += row_h + air


# Scale


def scale(slide, field: Field, body: Body, f: Palette, typeface: str) -> None:
    """A bar from light to full colour between two named ends, each label above its place on it."""
    typeface, g, left, top, width, height = _start(field, typeface)
    marks: list[Mark] = body.marks
    label_w = min(width / max(1, len(marks)) * 1.8, width * 0.3)

    def measure(k: float):
        text_pt = g * k
        title_pt = grid.TITLE_RATIO * g * k
        label_h = max(_h(m.label, text_pt, label_w, typeface, True) for m in marks)
        end_h = max(_h(t, title_pt, width * 0.45, typeface, True) for t in (body.left, body.right))
        bar = 1.1 * g * k
        marker = 0.8 * g * k
        h = 2 * label_h + 0.4 * g * k + marker + 0.3 * g * k + bar + 0.6 * g * k + end_h
        whole = (_fits([m.label for m in marks], typeface, text_pt, label_w, True)
                 and _fits([body.left, body.right], typeface, title_pt, width * 0.45, True))
        return whole and h <= height, (text_pt, title_pt, label_h, end_h, bar, marker, h, k)

    text_pt, title_pt, label_h, end_h, bar, marker, h, k = _largest(measure, g, 1.0)[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    bar_y = top + 2 * label_h + 0.4 * g * k + marker + 0.3 * g * k
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(left), _emu(bar_y), _emu(width), _emu(bar))
    shape.adjustments[0] = 0.5
    shape.fill.gradient()
    shape.fill.gradient_angle = 0
    stops = shape.fill.gradient_stops
    stops[0].color.rgb, stops[0].position = f.tint(fb.AREA), 0.0
    stops[1].color.rgb, stops[1].position = f.tint(fb.FULL), 1.0
    shape.line.fill.background()
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    for i, m in enumerate(marks):
        x = left + m.pos * width
        label_x = min(max(left, x - label_w / 2), left + width - label_w)
        # Neighbours alternate between two rows, so close places do not print over each other.
        y = top + (label_h if i % 2 == 0 else 0.0)
        _text(slide, label_x, y, label_w, label_h, m.label, text_pt, _BLACK, True, typeface,
              anchor=MSO_ANCHOR.BOTTOM)
        tip = bar_y - 0.15 * g * k
        _polygon(slide, [(x - marker / 2, tip - marker), (x + marker / 2, tip - marker), (x, tip)], f.tint(fb.FULL))
    end_y = bar_y + bar + 0.6 * g * k
    _text(slide, left, end_y, width * 0.45, end_h, body.left, title_pt, f.text_color(), True, typeface,
          align=PP_ALIGN.LEFT)
    _text(slide, left + width * 0.55, end_y, width * 0.45, end_h, body.right, title_pt, f.text_color(), True,
          typeface, align=PP_ALIGN.RIGHT)


# Venn


def _see_through(shape, alpha: int) -> None:
    color = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    etree.SubElement(color, qn("a:alpha")).set("val", str(alpha))


def venn(slide, field: Field, sides: list[Side], shared: list[str], f: Palette, typeface: str) -> None:
    """Two overlapping circles, what only one has inside it, what both share where they overlap."""
    typeface, g, left, top, width, height = _start(field, typeface)
    d = min(height * _VENN_HEIGHT, width / (2 - _VENN_OVERLAP))
    overlap = d * _VENN_OVERLAP
    x0 = left + (width - (2 * d - overlap)) / 2
    y0 = top + max(0.0, (height - d) * _TOP_SHARE)
    own_w = (d - overlap) * 0.82
    lens_w = overlap * 0.9
    room = d * 0.62

    def block(side: Side, title_pt: float, text_pt: float) -> float:
        return (_h(side.title, title_pt, own_w, typeface, True) + 0.4 * text_pt
                + sum(_h(t, text_pt, own_w, typeface) + 0.3 * text_pt for t in side.items))

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        both = sum(_h(t, text_pt, lens_w, typeface, True) + 0.3 * text_pt for t in shared)
        h = max(max(block(s, title_pt, text_pt) for s in sides), both)
        whole = (_fits([s.title for s in sides], typeface, title_pt, own_w, True)
                 and _fits([t for s in sides for t in s.items], typeface, text_pt, own_w)
                 and _fits(shared, typeface, text_pt, lens_w, True))
        return whole and h <= room, (title_pt, text_pt, both)

    title_pt, text_pt, both = _largest(measure, g)[2]
    for i, ratio in enumerate((0.55, 0.35)):
        circle = _shape(slide, MSO_SHAPE.OVAL, x0 + i * (d - overlap), y0, d, d, f.tint(ratio))
        _see_through(circle, 60000)
    cy = y0 + d / 2
    for i, side in enumerate(sides):
        x = x0 + (d - overlap - own_w) / 2 + 0.05 * d if i == 0 else x0 + d + (d - overlap - own_w) / 2 - 0.05 * d
        y = cy - block(side, title_pt, text_pt) / 2
        title_h = _h(side.title, title_pt, own_w, typeface, True)
        _text(slide, x, y, own_w, title_h, side.title, title_pt, _BLACK, True, typeface)
        y += title_h + 0.4 * text_pt
        for t in side.items:
            item_h = _h(t, text_pt, own_w, typeface)
            _text(slide, x, y, own_w, item_h, t, text_pt, _BLACK, False, typeface)
            y += item_h + 0.3 * text_pt
    y = cy - both / 2
    for t in shared:
        item_h = _h(t, text_pt, lens_w, typeface, True)
        _text(slide, x0 + d - overlap / 2 - lens_w / 2, y, lens_w, item_h, t, text_pt, _BLACK, True, typeface)
        y += item_h + 0.3 * text_pt


# Tree


def tree(slide, field: Field, root: str, branches: list[Branch], f: Palette, typeface: str) -> None:
    """The whole on top in the colour, its parts in a row below it joined by lines, their subparts listed under them."""
    typeface, g, left, top, width, height = _start(field, typeface)
    n = len(branches)
    gap = grid.GAP * g * 1.5
    col_w = (width - (n - 1) * gap) / n
    root_w = min(width * 0.45, max(col_w * 1.2, width * 0.3))

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        pad = 0.6 * g * k
        root_h = _h(root, title_pt * 1.1, root_w - 2 * pad, typeface, True) + 2 * pad
        branch_h = max(_h(b.title, title_pt, col_w - 2 * pad, typeface, True) for b in branches) + 2 * pad
        dot = 0.35 * text_pt
        leaf_w = col_w - dot - 0.5 * text_pt
        leaves = max(sum(_h(t, text_pt, leaf_w, typeface) + 0.35 * text_pt for t in b.leaves) for b in branches)
        join = 1.6 * g * k
        h = root_h + join + branch_h + 0.6 * g * k + leaves
        whole = (_fits([root], typeface, title_pt * 1.1, root_w - 2 * pad, True)
                 and _fits([b.title for b in branches], typeface, title_pt, col_w - 2 * pad, True)
                 and _fits([t for b in branches for t in b.leaves], typeface, text_pt, leaf_w))
        return whole and h <= height, (title_pt, text_pt, pad, root_h, branch_h, dot, leaf_w, join, h, k)

    title_pt, text_pt, pad, root_h, branch_h, dot, leaf_w, join, h, k = _largest(measure, g)[2]
    top += max(0.0, (height - h) * _TOP_SHARE)
    cx = left + width / 2
    _panel(slide, cx - root_w / 2, top, root_w, root_h, f.tint(fb.FULL))
    _text(slide, cx - root_w / 2 + pad, top, root_w - 2 * pad, root_h, root, title_pt * 1.1, f.text_on(), True,
          typeface, anchor=MSO_ANCHOR.MIDDLE)
    centres = [left + i * (col_w + gap) + col_w / 2 for i in range(n)]
    mid = top + root_h + join / 2
    color = f.tint(fb.LINE)
    _shape(slide, MSO_SHAPE.RECTANGLE, cx - _LINE_PT / 2, top + root_h, _LINE_PT, join / 2, color)
    _shape(slide, MSO_SHAPE.RECTANGLE, centres[0], mid - _LINE_PT / 2, centres[-1] - centres[0], _LINE_PT, color)
    branch_y = top + root_h + join
    for x_c, b in zip(centres, branches):
        _shape(slide, MSO_SHAPE.RECTANGLE, x_c - _LINE_PT / 2, mid, _LINE_PT, join / 2, color)
        x = x_c - col_w / 2
        _panel(slide, x, branch_y, col_w, branch_h, f.tint(fb.AREA))
        _text(slide, x + pad, branch_y, col_w - 2 * pad, branch_h, b.title, title_pt, f.text_color(), True,
              typeface, anchor=MSO_ANCHOR.MIDDLE)
        y = branch_y + branch_h + 0.6 * g * k
        for t in b.leaves:
            leaf_h = _h(t, text_pt, leaf_w, typeface)
            _shape(slide, MSO_SHAPE.OVAL, x + 0.1 * text_pt, y + (text_pt * LINE_RATIO - dot) / 2, dot, dot,
                   f.tint(fb.FULL))
            _text(slide, x + dot + 0.5 * text_pt, y, leaf_w, leaf_h, t, text_pt, _BLACK, False, typeface,
                  align=PP_ALIGN.LEFT)
            y += leaf_h + 0.35 * text_pt


# Worked example


def example(slide, field: Field, body: Body, f: Palette, typeface: str) -> None:
    """The situation in a light panel, the steps as numbered boxes with arrows between them, the result in the colour."""
    typeface, g, left, top, width, height = _start(field, typeface)
    steps = body.items
    n = len(steps)
    arrow = 1.6 * g
    step_w = (width - (n - 1) * arrow) / n

    def measure(k: float):
        title_pt, text_pt = _sizes(g, k)
        pad = 0.7 * g * k
        bar = 0.3 * g
        situation_h = _h(body.situation, text_pt, width - 2 * pad - bar, typeface) + 2 * pad if body.situation else 0.0
        badge = 1.2 * text_pt
        step_h = max(_h(t, text_pt, step_w - 2 * pad, typeface) for t in steps) + 2 * pad + badge + 0.4 * text_pt
        result_h = _h(body.result, title_pt, width - 2 * pad, typeface, True) + 2 * pad
        air = 0.8 * g * k
        h = situation_h + (air if situation_h else 0.0) + step_h + air + result_h
        whole = (_fits([body.situation, *steps], typeface, text_pt, min(width - 2 * pad - bar, step_w - 2 * pad))
                 and _fits([body.result], typeface, title_pt, width - 2 * pad, True))
        return whole and h <= height, (title_pt, text_pt, pad, bar, situation_h, badge, step_h, result_h, air, h)

    title_pt, text_pt, pad, bar, situation_h, badge, step_h, result_h, air, h = _largest(measure, g)[2]
    y = top + max(0.0, (height - h) * _TOP_SHARE)
    if situation_h:
        _panel(slide, left, y, width, situation_h, f.tint(fb.AREA), round_=False)
        _shape(slide, MSO_SHAPE.RECTANGLE, left, y, bar, situation_h, f.tint(fb.FULL))
        _text(slide, left + bar + pad, y, width - 2 * pad - bar, situation_h, body.situation, text_pt, _BLACK, False,
              typeface, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        y += situation_h + air
    for i, t in enumerate(steps):
        x = left + i * (step_w + arrow)
        _panel(slide, x, y, step_w, step_h, f.tint(fb.BAND))
        _number(slide, x + pad, y + pad, badge, i + 1, f, typeface)
        _text(slide, x + pad, y + pad + badge + 0.4 * text_pt, step_w - 2 * pad, step_h - 2 * pad - badge, t, text_pt,
              _BLACK, False, typeface, align=PP_ALIGN.LEFT)
        if i < n - 1:
            size = arrow * 0.6
            _shape(slide, MSO_SHAPE.RIGHT_ARROW, x + step_w + (arrow - size) / 2, y + step_h / 2 - size * 0.35,
                   size, size * 0.7, f.tint(fb.LINE))
    y += step_h + air
    _panel(slide, left, y, width, result_h, f.tint(fb.FULL))
    _text(slide, left + pad, y, width - 2 * pad, result_h, body.result, title_pt, f.text_on(), True, typeface,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)


def draw(slide, field: Field, body: Body, f: Palette, typeface: str) -> bool:
    """Draws a body of one of these forms, False when the body is of another kind or empty."""
    k = body.kind
    if k == "timeline" and body.events:
        timeline(slide, field, body.events, f, typeface)
    elif k == "cycle" and body.cycle:
        cycle(slide, field, body.cycle, body.center, f, typeface)
    elif k == "versus" and len(body.sides) == 2:
        versus(slide, field, body.sides, body.tone, f, typeface)
    elif k == "matrix" and body.cells:
        matrix(slide, field, body.cells, body.axis_x, body.axis_y, f, typeface)
    elif k == "quote" and body.text:
        quote(slide, field, body, f, typeface)
    elif k == "checklist" and body.items:
        checklist(slide, field, body.items, f, typeface)
    elif k == "definitions" and body.terms:
        definitions(slide, field, body.terms, f, typeface)
    elif k == "mistakes" and body.pairs:
        mistakes(slide, field, body.pairs, f, typeface)
    elif k == "scale" and body.marks:
        scale(slide, field, body, f, typeface)
    elif k == "venn" and len(body.sides) == 2:
        venn(slide, field, body.sides, body.shared, f, typeface)
    elif k == "tree" and body.branches:
        tree(slide, field, body.root, body.branches, f, typeface)
    elif k == "example" and body.items and body.result:
        example(slide, field, body, f, typeface)
    else:
        return False
    return True
