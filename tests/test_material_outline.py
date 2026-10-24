"""The outline over documents: what the model is asked, and what reaches the deck.

The documents are real files, built by the fixtures and read by the reader. The fake model
echoes what it was sent, so the assertions are about the material that actually arrived and
about the text that actually stands in the finished .pptx.
"""

import json
from io import BytesIO

import pytest
from pptx import Presentation

from conftest import PDF_FOOTER, slides_answer
from outline2pptx import api, budget as bd, material as mt
from outline2pptx import outline as ol
from outline2pptx.errors import LimitError, NoReadableText


def options(server, **extra):
    return {"base_url": server.base_url, "model": "a-model", **extra}


def material_lines(body, count=5):
    """The first lines of the material the server was sent, as an outline answer."""
    user = body["messages"][1]["content"]
    block = user.split("MATERIAL:\n", 1)[1]
    lines = [line.strip() for line in block.splitlines()
             if line.strip() and not line.startswith("#")][:count]
    roles = ["cover"] + ["text"] * (len(lines) - 2) + ["closing"]
    return json.dumps({"title": lines[0], "subtitle": "From the handout",
                       "slides": [{"role": roles[i], "title": line[:110], "gist": line[:110]}
                                  for i, line in enumerate(lines)]})


def echo_slides(body):
    """A content answer that writes each row's own title and gist onto its slide."""
    user = body["messages"][1]["content"]
    rows = []
    for line in user.splitlines():
        if line.startswith("SLIDE "):
            rows.append((line.split()[1], line.split("gist: ", 1)[-1]))
    answer = slides_answer([sid for sid, _ in rows])
    for entry, (_, gist) in zip(answer["slides"], rows):
        entry["title"] = gist[:110]
        entry["body"] = {"kind": "bullets", "items": [gist[:110]] * 5}
        entry["columns"] = [{"heading": "Part", "items": [gist[:110]] * 4} for _ in range(2)]
    return answer


# --- What the outline call is asked ------------------------------------------------------

def test_the_call_over_documents_carries_the_material_and_not_a_topic(server, handout_pdf):
    server.outline = material_lines
    api.outline_from_material([handout_pdf], **options(server))
    sent = server.requests[0]["body"]["messages"]
    system, user = sent[0]["content"], sent[1]["content"]
    assert "From the material you pick what matters" in system
    assert "Work through the material to its last line" in system
    assert "MATERIAL:\n" in user
    assert "TOPIC:" not in user
    # The body of the handout is there, its repeating footer is not.
    assert "Jahresarbeitszahl" in user
    assert PDF_FOOTER not in user
    assert "# leitfaden.pdf" in user


def test_a_single_file_is_not_asked_to_be_balanced_against_others(server, handout_docx):
    server.outline = material_lines
    api.outline_from_material(handout_docx, **options(server))
    system = server.requests[0]["body"]["messages"][0]["content"]
    assert "EVERY one of them gets slides of its own" not in system


def test_several_files_each_get_their_own_slides_and_a_share(server, handout_pdf, handout_docx):
    server.outline = material_lines
    api.outline_from_material([handout_pdf, handout_docx], slides=12, **options(server))
    system = server.requests[0]["body"]["messages"][0]["content"]
    assert "comes from 2 files: leitfaden.pdf, leitfaden.docx" in system
    assert "EVERY one of them gets slides of its own" in system
    # 12 slides less cover, overview and closing, shared out by how much each file brings.
    shares = ol.slides_per_file(mt.read([handout_pdf, handout_docx]).weights, 12 - ol.FRAME_SLIDES)
    assert sum(shares.values()) == 9
    for name, number in shares.items():
        assert f"{name}: {number}" in system


def test_the_amount_of_material_decides_the_length_of_the_deck(server, handout_pdf):
    server.outline = material_lines
    api.outline_from_material([handout_pdf], **options(server))
    wanted = ol.target_slides(mt.read([handout_pdf]).chars)
    assert f"Plan around {wanted} slides" in server.requests[0]["body"]["messages"][0]["content"]
    # A number that was asked for wins over the one the material suggests.
    api.outline_from_material([handout_pdf], slides=17, **options(server))
    assert "Plan around 17 slides" in server.requests[1]["body"]["messages"][0]["content"]


def test_the_target_grows_with_the_material_and_stays_inside_the_bounds():
    lo, hi = ol.SLIDE_BOUNDS
    assert ol.target_slides(0) == 8
    assert ol.target_slides(200_000) == 11
    assert ol.target_slides(10_000_000) == hi
    assert ol.target_slides(-5) >= lo


def test_the_slides_of_a_file_follow_its_size_but_never_run_away_with_the_deck():
    share = ol.slides_per_file({"big.pdf": 90_000, "small.pdf": 10_000}, 9)
    assert sum(share.values()) == 9
    assert share["big.pdf"] > share["small.pdf"]
    assert share["big.pdf"] <= 2 * share["small.pdf"]
    # One file shares nothing, and fewer slides than files shares nothing either.
    assert ol.slides_per_file({"only.pdf": 500}, 9) == {}
    assert ol.slides_per_file({"a": 1, "b": 1, "c": 1}, 2) == {}


def test_an_outline_over_documents_is_asked_again_when_it_comes_back_unusable(server,
                                                                             handout_docx):
    first = True

    def reply(body):
        nonlocal first
        if first:
            first = False
            return "no json at all"
        return material_lines(body)

    server.outline = reply
    text = api.outline_from_material([handout_docx], **options(server))
    assert text.startswith("# ")
    assert server.kinds() == ["outline", "outline"]


# --- The size budget in the outline step -------------------------------------------------

def test_material_larger_than_the_window_is_mapped_first_and_planned_from_its_topics(
        server, tmp_path):
    limit = bd.window_chars(32_768)
    path = tmp_path / "lang.txt"
    # Deliberately larger than the window, a short text would never reach the budget.
    path.write_text("\n\n".join(f"Abschnitt {n}: Die Jahresarbeitszahl steigt mit der Quelle. "
                                + "Fuellsatz. " * 200 for n in range(1, 40)))
    assert path.stat().st_size > limit

    server.topics = lambda body: json.dumps({"topics": [
        {"title": "Quelle und Zahl", "gist": "Was das Material sagt."},
        {"title": "Vorlauftemperatur", "gist": "Welche Grade genannt werden."},
        {"title": "Heizstab", "gist": "Wann er anspringt."}]})
    server.outline = material_lines
    api.outline_from_material([path], context_window=32_768, **options(server))

    kinds = server.kinds()
    assert kinds[-1] == "outline"
    assert kinds.count("topics") == len(bd.pieces(mt.read([path]).block, limit)) > 1
    planned = server.requests[-1]["body"]["messages"][1]["content"]
    assert "1. Quelle und Zahl: Was das Material sagt." in planned
    # The call stays inside the window instead of being cut off at the end by the server.
    assert len(planned.split("MATERIAL:\n", 1)[1]) <= limit


def test_material_that_fits_the_window_costs_no_mapping_call(server, handout_pdf):
    server.outline = material_lines
    api.outline_from_material([handout_pdf], **options(server))
    assert server.kinds() == ["outline"]


def test_a_window_too_small_for_a_deck_is_refused_before_any_call(server, handout_pdf):
    from outline2pptx.errors import ConfigError

    with pytest.raises(ConfigError, match="--context-window"):
        api.outline_from_material([handout_pdf], context_window=1000, **options(server))
    assert server.requests == []


def test_material_over_the_character_limit_never_reaches_the_model(server, tmp_path):
    path = tmp_path / "lang.txt"
    path.write_text("Ein Satz ueber die Waermepumpe. " * 400)
    with pytest.raises(LimitError):
        api.outline_from_material([path], max_material_chars=1000, **options(server))
    assert server.requests == []


def test_a_document_without_text_is_refused_before_any_call(server, tmp_path):
    path = tmp_path / "leer.txt"
    path.write_text("   \n  \n")
    with pytest.raises(NoReadableText):
        api.outline_from_material([path], **options(server))
    assert server.requests == []


# --- From the document to the finished file ----------------------------------------------

def test_a_deck_built_from_a_handout_carries_its_body_and_not_its_page_footer(
        server, handout_pdf, tmp_path):
    server.outline = material_lines
    server.slides = echo_slides
    outline_md = api.outline_from_material([handout_pdf], **options(server))
    assert "Kapitel 1: Die Waermequelle" in outline_md

    out = tmp_path / "deck.pptx"
    deck = api.build(outline_md, out, **options(server))
    assert out.exists()

    prs = Presentation(BytesIO(out.read_bytes()))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    whole = "\n".join(texts)
    # What the handout says stands on the slides.
    assert "Waermequelle" in whole
    assert "Jahresarbeitszahl" in whole
    # What only its footer said stands nowhere.
    assert PDF_FOOTER not in whole
    assert "Beispielverlag" not in whole
    assert "Seite 3 von 6" not in whole
    assert len(prs.slides._sldIdLst) == len(deck.slides)
