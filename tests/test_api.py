import json
from io import BytesIO

import pytest
from pptx import Presentation

from conftest import outline_answer, slides_answer
from outline2pptx import api
from outline2pptx.errors import ConfigError, ModelError

OUTLINE_MD = """
# Bees and Flowers
How pollination really works

## What bees do all day

## Pollination step by step
Pollen travels from one flower to the next.

## [two_columns] Winners and losers
Who gains and who pays.

## [text_with_image] A hive from the inside
What the comb looks like up close.

## Plant one thing this spring
Small gardens carry whole colonies.
"""


def options(server, **extra):
    return {"base_url": server.base_url, "model": "a-model", **extra}


def answer_for_batch(body):
    """Answer a content call with exactly the slide ids it asked for."""
    user = body["messages"][1]["content"]
    ids = [line.split()[1] for line in user.splitlines() if line.startswith("SLIDE ")]
    return slides_answer(ids)


def test_outline_from_topic_returns_markdown(server):
    server.outline = outline_answer(6, title="Bees")
    text = api.outline_from_topic("Bees", slides=6, **options(server))
    assert text.startswith("# Bees")
    assert text.count("## ") == 5
    assert server.kinds() == ["outline"]
    assert "around 6 slides" in server.requests[0]["body"]["messages"][0]["content"]


def test_outline_is_asked_again_when_it_is_unusable(server):
    answers = iter(["not json at all", json.dumps(outline_answer(5))])
    server.outline = lambda body: next(answers)
    text = api.outline_from_topic("Bees", **options(server))
    assert text.startswith("# ")
    assert server.kinds() == ["outline", "outline"]


def test_outline_gives_up_after_the_retries(server):
    server.outline = "still not json"
    with pytest.raises(ModelError):
        api.outline_from_topic("Bees", **options(server))
    assert len(server.requests) == api.OUTLINE_TRIES


def test_outline_needs_a_topic(server):
    with pytest.raises(ModelError, match="No topic"):
        api.outline_from_topic("  ", **options(server))


def test_build_writes_a_deck(server, tmp_path):
    server.slides = answer_for_batch
    out = tmp_path / "deck.pptx"
    deck = api.build(OUTLINE_MD, out, **options(server))
    assert out.exists()
    assert deck.info["slides"] == len(deck.outline.rows)
    assert deck.info["template"] == "Clean"
    assert deck.info["model"] == "a-model"
    assert deck.info["image_model"] == ""
    prs = Presentation(BytesIO(deck.data))
    assert len(prs.slides._sldIdLst) == len(deck.slides)
    assert all(k == "slides" for k in server.kinds())


def test_build_batches_the_slides(server):
    server.slides = answer_for_batch
    api.build(OUTLINE_MD, batch=2, **options(server))
    # six rows in batches of two
    assert len(server.requests) == 3


def test_build_asks_again_for_a_thin_slide(server):
    seen = {"n": 0}

    def reply(body):
        seen["n"] += 1
        user = body["messages"][1]["content"]
        ids = [line.split()[1] for line in user.splitlines() if line.startswith("SLIDE ")]
        if seen["n"] == 1:
            thin = slides_answer(ids)
            for entry in thin["slides"]:
                entry["body"] = {"kind": "bullets", "items": ["Only one"]}
            return thin
        return slides_answer(ids)

    server.slides = reply
    deck = api.build(OUTLINE_MD, batch=10, **options(server))
    assert len(server.requests) == 2
    assert "too thin" in server.requests[1]["body"]["messages"][-1]["content"]
    bullets = [s for s in deck.slides if s.body is not None and s.body.kind == "bullets"]
    assert bullets and all(len(s.body.items) > 1 for s in bullets)


def test_build_keeps_the_first_version_when_the_second_ask_fails(server):
    state = {"n": 0}

    def reply(body):
        state["n"] += 1
        user = body["messages"][1]["content"]
        ids = [line.split()[1] for line in user.splitlines() if line.startswith("SLIDE ")]
        if state["n"] == 1:
            thin = slides_answer(ids)
            for entry in thin["slides"]:
                entry["body"] = {"kind": "bullets", "items": ["Only one"]}
            return thin
        return "nonsense"

    server.slides = reply
    deck = api.build(OUTLINE_MD, batch=10, **options(server))
    assert deck.info["slides"] == 6


def test_build_gives_up_when_every_answer_is_unusable(server):
    server.slides = "nonsense"
    with pytest.raises(ModelError):
        api.build(OUTLINE_MD, batch=10, **options(server))
    assert len(server.requests) == api.SLIDE_TRIES


def test_build_without_a_base_url_says_what_is_missing():
    with pytest.raises(ConfigError, match="base_url"):
        api.build(OUTLINE_MD)


def test_build_with_an_unknown_template_names_the_known_ones(server):
    from outline2pptx.errors import TemplateError

    with pytest.raises(TemplateError, match="clean"):
        api.build(OUTLINE_MD, template="nope", **options(server))


def test_build_with_notes_puts_them_in_the_file(server):
    def reply(body):
        user = body["messages"][1]["content"]
        ids = [line.split()[1] for line in user.splitlines() if line.startswith("SLIDE ")]
        return slides_answer(ids, notes="This is what the speaker says about the slide.")

    server.slides = reply
    deck = api.build(OUTLINE_MD, notes=True, **options(server))
    prs = Presentation(BytesIO(deck.data))
    assert any(s.has_notes_slide for s in prs.slides)


def test_build_with_images_calls_the_image_model(server):
    server.slides = answer_for_batch
    deck = api.build(OUTLINE_MD, images=True, image_model="draw", **options(server))
    assert "image" in server.kinds()
    assert deck.info["image_model"] == "draw"


def test_build_reports_progress(server):
    server.slides = answer_for_batch
    seen = []
    api.build(OUTLINE_MD, progress=lambda stage, done, total: seen.append(stage), **options(server))
    assert "slides" in seen and "pictures" in seen and seen[-1] == "done"


def test_list_templates_names_the_built_in_ones():
    ids = [t["id"] for t in api.list_templates()]
    assert ids == ["clean", "forest", "warm"]
    assert all("path" not in t for t in api.list_templates())


def test_a_template_dir_of_your_own_comes_first(tmp_path):
    from outline2pptx import library

    source = library.BUILTIN / "clean.pptx"
    (tmp_path / "mine.pptx").write_bytes(source.read_bytes())
    ids = [t["id"] for t in api.list_templates(tmp_path)]
    assert ids[0] == "mine"


def test_json_mode_is_dropped_when_the_provider_refuses_it(server):
    server.reject_json_mode = True
    server.outline = outline_answer(5)
    api.outline_from_topic("Bees", **options(server))
    assert len(server.requests) == 2
    assert "response_format" in server.requests[0]["body"]
    assert "response_format" not in server.requests[1]["body"]


def test_the_api_key_reaches_the_server(server):
    server.outline = outline_answer(5)
    api.outline_from_topic("Bees", api_key="secret-key", **options(server))
    assert server.requests[0]["headers"]["authorization"] == "Bearer secret-key"


def test_without_a_key_a_placeholder_is_sent(server):
    from outline2pptx.config import PLACEHOLDER_API_KEY

    server.outline = outline_answer(5)
    api.outline_from_topic("Bees", **options(server))
    assert server.requests[0]["headers"]["authorization"] == f"Bearer {PLACEHOLDER_API_KEY}"
