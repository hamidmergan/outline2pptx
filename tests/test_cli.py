import io
import json
from io import BytesIO

import pytest
from pptx import Presentation

from conftest import PDF_FOOTER, outline_answer, slides_answer
from outline2pptx import cli
from test_api import OUTLINE_MD, answer_for_batch
from test_material_outline import echo_slides, material_lines


def flags(server):
    return ["--base-url", server.base_url, "--model", "a-model"]


@pytest.fixture
def outline_file(tmp_path):
    path = tmp_path / "outline.md"
    path.write_text(OUTLINE_MD, "utf8")
    return path


def test_help_is_english(capsys):
    with pytest.raises(SystemExit):
        cli.main(["--help"])
    text = capsys.readouterr().out
    assert "Build a PowerPoint deck" in text
    assert "--base-url" in text and "--template" in text


def test_version_prints_the_number(capsys):
    with pytest.raises(SystemExit):
        cli.main(["--version"])
    assert "0.1.0" in capsys.readouterr().out


def test_list_templates_shows_the_built_in_ones(capsys):
    assert cli.main(["--list-templates"]) == 0
    out = capsys.readouterr().out
    assert "clean" in out and "forest" in out and "warm" in out


def test_build_from_a_file_writes_the_deck(server, outline_file, tmp_path, capsys):
    server.slides = answer_for_batch
    out = tmp_path / "deck.pptx"
    assert cli.main([str(outline_file), "-o", str(out), *flags(server)]) == 0
    assert out.exists()
    prs = Presentation(BytesIO(out.read_bytes()))
    assert len(prs.slides._sldIdLst) >= 6
    assert "Wrote" in capsys.readouterr().err


def test_the_output_defaults_to_the_outline_name(server, outline_file):
    server.slides = answer_for_batch
    assert cli.main([str(outline_file), *flags(server)]) == 0
    assert outline_file.with_suffix(".pptx").exists()


def test_an_outline_can_come_from_standard_input(server, tmp_path, monkeypatch):
    server.slides = answer_for_batch
    monkeypatch.setattr("sys.stdin", io.StringIO(OUTLINE_MD))
    out = tmp_path / "deck.pptx"
    assert cli.main(["-", "-o", str(out), *flags(server)]) == 0
    assert out.exists()


def test_a_topic_plans_the_outline_first(server, tmp_path):
    server.outline = outline_answer(6, title="Bees")
    server.slides = answer_for_batch
    out = tmp_path / "deck.pptx"
    assert cli.main(["--topic", "Bees", "--slides", "6", "-o", str(out), *flags(server)]) == 0
    assert server.kinds()[0] == "outline"
    assert out.exists()


def test_outline_only_prints_markdown_and_builds_nothing(server, capsys, tmp_path):
    server.outline = outline_answer(5, title="Bees")
    out = tmp_path / "deck.pptx"
    assert cli.main(["--topic", "Bees", "--outline-only", "-o", str(out), *flags(server)]) == 0
    assert capsys.readouterr().out.startswith("# Bees")
    assert not out.exists()


def test_images_switch_on_the_image_model(server, outline_file, tmp_path):
    server.slides = answer_for_batch
    out = tmp_path / "deck.pptx"
    code = cli.main([str(outline_file), "-o", str(out), "--images", "--image-model", "draw",
                     *flags(server)])
    assert code == 0
    assert "image" in server.kinds()


def test_notes_reach_the_file(server, outline_file, tmp_path):
    def reply(body):
        user = body["messages"][1]["content"]
        ids = [line.split()[1] for line in user.splitlines() if line.startswith("SLIDE ")]
        return slides_answer(ids, notes="What the speaker says about this slide.")

    server.slides = reply
    out = tmp_path / "deck.pptx"
    assert cli.main([str(outline_file), "-o", str(out), "--notes", *flags(server)]) == 0
    prs = Presentation(BytesIO(out.read_bytes()))
    assert any(s.has_notes_slide for s in prs.slides)


def test_without_an_outline_a_topic_or_material_it_asks_for_one(capsys):
    assert cli.main([]) == 2
    assert "Give an outline file, --topic or --material." in capsys.readouterr().err


def test_two_of_the_three_inputs_together_are_refused(capsys, outline_file, handout_pdf):
    assert cli.main([str(outline_file), "--topic", "Bees"]) == 2
    assert "not an outline file and --topic" in capsys.readouterr().err
    assert cli.main(["--topic", "Bees", "--material", str(handout_pdf)]) == 2
    assert "not --topic and --material" in capsys.readouterr().err


def test_a_missing_outline_file_is_reported(capsys, tmp_path):
    assert cli.main([str(tmp_path / "nope.md"), "--base-url", "http://x/v1", "--model", "m"]) == 2
    assert "No such outline file" in capsys.readouterr().err


def test_a_missing_base_url_names_the_flag(capsys, outline_file):
    assert cli.main([str(outline_file)]) == 2
    err = capsys.readouterr().err
    assert "--base-url" in err and "--model" in err


def test_an_unknown_template_names_the_known_ones(capsys, outline_file, server):
    assert cli.main([str(outline_file), "--template", "nope", *flags(server)]) == 2
    assert "clean" in capsys.readouterr().err


def test_images_without_an_image_model_are_refused(capsys, outline_file, server):
    assert cli.main([str(outline_file), "--images", *flags(server)]) == 2
    assert "image_model" in capsys.readouterr().err


def test_settings_can_come_from_the_environment(server, outline_file, tmp_path, monkeypatch):
    server.slides = answer_for_batch
    monkeypatch.setenv("OUTLINE2PPTX_BASE_URL", server.base_url)
    monkeypatch.setenv("OUTLINE2PPTX_MODEL", "a-model")
    out = tmp_path / "deck.pptx"
    assert cli.main([str(outline_file), "-o", str(out)]) == 0
    assert out.exists()


def test_a_failed_preview_keeps_the_deck_it_already_wrote(server, outline_file, tmp_path,
                                                          monkeypatch, capsys):
    server.slides = answer_for_batch
    monkeypatch.setattr("outline2pptx.preview.available", lambda: False)
    out = tmp_path / "deck.pptx"
    code = cli.main([str(outline_file), "-o", str(out), "--preview", str(tmp_path / "png"),
                     *flags(server)])
    err = capsys.readouterr().err
    assert code == 2
    assert "LibreOffice" in err
    # The deck is the point of the run. It stays on disk and the message says which half failed.
    assert out.exists() and out.stat().st_size > 0
    assert "only the preview is missing" in err
    assert Presentation(BytesIO(out.read_bytes())).slides


def test_the_preview_writes_one_png_per_slide_of_the_deck(server, outline_file, tmp_path,
                                                          monkeypatch):
    server.slides = answer_for_batch
    monkeypatch.setattr("outline2pptx.render.available", lambda: True)

    def one_per_slide(data, dpi=80):
        count = len(Presentation(BytesIO(data)).slides._sldIdLst)
        return [f"page {n}".encode() for n in range(1, count + 1)]

    monkeypatch.setattr("outline2pptx.preview.render", one_per_slide)
    out = tmp_path / "deck.pptx"
    folder = tmp_path / "png"
    assert cli.main([str(outline_file), "-o", str(out), "--preview", str(folder),
                     *flags(server)]) == 0
    slides = len(Presentation(BytesIO(out.read_bytes())).slides._sldIdLst)
    files = sorted(folder.iterdir())
    assert [f.name for f in files] == [f"slide-{n:02d}.png" for n in range(1, slides + 1)]
    assert [f.read_bytes() for f in files] == [f"page {n}".encode()
                                               for n in range(1, slides + 1)]


def test_slides_with_an_outline_file_is_refused(capsys, outline_file, server):
    code = cli.main([str(outline_file), "--slides", "12", *flags(server)])
    assert code == 2
    assert "--slides" in capsys.readouterr().err


def test_serve_without_the_extra_packages_says_what_to_install(capsys, monkeypatch):
    import builtins

    real = builtins.__import__

    def no_uvicorn(name, *args, **kwargs):
        if name == "uvicorn":
            raise ImportError("no uvicorn")
        return real(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_uvicorn)
    assert cli.main(["serve"]) == 2
    assert "outline2pptx[ui]" in capsys.readouterr().err


# --- Documents on the command line -------------------------------------------------------

def test_material_plans_the_outline_over_the_document_and_builds_the_deck(server, handout_pdf,
                                                                         tmp_path, capsys):
    server.outline = material_lines
    server.slides = echo_slides
    out = tmp_path / "deck.pptx"
    assert cli.main(["--material", str(handout_pdf), "-o", str(out), *flags(server)]) == 0
    assert out.exists()
    err = capsys.readouterr().err
    # The size of the material is on screen, so the budget is no surprise.
    assert "Material: 1 file(s), 6 page(s)," in err
    assert "characters." in err
    assert server.kinds()[0] == "outline"
    prs = Presentation(BytesIO(out.read_bytes()))
    texts = [shape.text_frame.text for slide in prs.slides for shape in slide.shapes
             if shape.has_text_frame]
    assert any("Waermequelle" in text for text in texts)
    assert all(PDF_FOOTER not in text for text in texts)


def test_several_documents_can_be_given_at_once(server, handout_pdf, handout_docx, tmp_path):
    server.outline = material_lines
    out = tmp_path / "deck.pptx"
    assert cli.main(["--material", str(handout_pdf), str(handout_docx), "-o", str(out),
                     "--outline-only", *flags(server)]) == 0
    system = server.requests[0]["body"]["messages"][0]["content"]
    assert "comes from 2 files: leitfaden.pdf, leitfaden.docx" in system


def test_the_output_defaults_to_the_name_of_the_first_document(server, handout_pdf, tmp_path,
                                                              monkeypatch):
    server.outline = material_lines
    server.slides = echo_slides
    monkeypatch.chdir(tmp_path)
    assert cli.main(["--material", str(handout_pdf), *flags(server)]) == 0
    assert (tmp_path / "leitfaden.pptx").exists()


def test_the_window_flag_decides_when_the_material_is_mapped_first(server, tmp_path, capsys):
    path = tmp_path / "lang.txt"
    path.write_text("\n\n".join(f"Abschnitt {n}: Die Jahresarbeitszahl steigt. "
                                 + "Fuellsatz. " * 200 for n in range(1, 40)))
    server.topics = lambda body: json.dumps({"topics": [
        {"title": "Quelle", "gist": "Was das Material sagt."},
        {"title": "Zahl", "gist": "Welche Zahlen es nennt."},
        {"title": "Heizstab", "gist": "Wann er anspringt."}]})
    server.outline = material_lines
    code = cli.main(["--material", str(path), "--context-window", "32768", "--outline-only",
                     *flags(server)])
    assert code == 0
    assert "planned through the topics of the material" in capsys.readouterr().err
    assert "topics" in server.kinds()


def test_a_document_the_deck_cannot_use_is_refused_by_name(server, tmp_path, capsys):
    path = tmp_path / "sound.mp3"
    path.write_bytes(b"not a document")
    assert cli.main(["--material", str(path), *flags(server)]) == 2
    err = capsys.readouterr().err
    assert ".mp3" in err and ".pdf" in err
    assert server.requests == []


def test_material_over_the_limit_names_the_flag(server, tmp_path, capsys):
    path = tmp_path / "lang.txt"
    path.write_text("Ein Satz ueber die Waermepumpe. " * 400)
    assert cli.main(["--material", str(path), "--max-material-chars", "1000",
                     *flags(server)]) == 2
    assert "--max-material-chars" in capsys.readouterr().err


def test_slides_may_be_asked_for_over_material(server, handout_pdf, capsys):
    server.outline = material_lines
    assert cli.main(["--material", str(handout_pdf), "--slides", "14", "--outline-only",
                     *flags(server)]) == 0
    assert "Plan around 14 slides" in server.requests[0]["body"]["messages"][0]["content"]
