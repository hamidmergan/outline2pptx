import base64
import io
import json
import os
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest


def _png(width=64, height=64, color=(120, 140, 200)):
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (width, height), color).save(out, format="PNG")
    return out.getvalue()


class FakeServer:
    """An OpenAI-compatible chat and image endpoint.

    ``outline`` answers the outline call, ``slides`` the content call. Both may be a dict, a raw
    string, or a callable that gets the request body.
    """

    def __init__(self):
        self.requests = []
        self.fail = False
        self.image_status = 200
        # How many image calls are refused before the first one is answered.
        self.refuse_images = 0
        self.reject_json_mode = False
        self.outline = {"title": "Deck", "subtitle": "", "slides": []}
        self.slides = {"slides": []}
        self.image_bytes = _png()
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_POST(self):
                length = int(self.headers.get("Content-Length", 0))
                body = json.loads(self.rfile.read(length) or b"{}")
                outer.requests.append({"path": self.path,
                                       "headers": {k.lower(): v for k, v in self.headers.items()},
                                       "body": body})
                if outer.fail:
                    return self.reply(500, b'{"error": "boom"}')
                if self.path.endswith("/images/generations"):
                    if outer.refuse_images > 0:
                        outer.refuse_images -= 1
                        return self.reply(400, b'{"error": {"message": "size"}}')
                    if outer.image_status != 200:
                        return self.reply(outer.image_status, b'{"error": {"message": "size"}}')
                    payload = {"created": 0, "data": [
                        {"b64_json": base64.b64encode(outer.image_bytes).decode()}]}
                    return self.reply(200, json.dumps(payload).encode())
                if outer.reject_json_mode and "response_format" in body:
                    return self.reply(400, b'{"error": {"message": "response_format not supported"}}')
                payload = {
                    "id": "x", "object": "chat.completion", "created": 0, "model": body["model"],
                    "choices": [{"index": 0, "finish_reason": "stop",
                                 "message": {"role": "assistant", "content": outer.answer(body)}}],
                }
                self.reply(200, json.dumps(payload).encode())

            def reply(self, status, data):
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

        self.httpd = HTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    def answer(self, body) -> str:
        system = body["messages"][0]["content"]
        source = self.outline if system.startswith("You plan slide decks") else self.slides
        if callable(source):
            source = source(body)
        return source if isinstance(source, str) else json.dumps(source)

    def kinds(self):
        """'outline', 'slides' or 'image' for every request so far."""
        out = []
        for r in self.requests:
            if r["path"].endswith("/images/generations"):
                out.append("image")
            elif r["body"]["messages"][0]["content"].startswith("You plan slide decks"):
                out.append("outline")
            else:
                out.append("slides")
        return out

    @property
    def base_url(self):
        return f"http://127.0.0.1:{self.httpd.server_port}/v1"

    def stop(self):
        self.httpd.shutdown()
        self.httpd.server_close()


@pytest.fixture
def server():
    s = FakeServer()
    yield s
    s.stop()


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    for name in list(os.environ):
        if name.startswith("OUTLINE2PPTX_") or name == "OPENAI_API_KEY":
            monkeypatch.delenv(name)


@pytest.fixture(autouse=True)
def fresh_json_mode():
    """Every test starts with JSON mode unprobed, the flag lives for the whole process."""
    from outline2pptx import llm

    llm._json_mode_supported = True
    yield
    llm._json_mode_supported = True


def outline_answer(count=5, title="A Deck", roles=None):
    """A well-formed outline answer with ``count`` slides."""
    roles = roles or ["cover"] + ["text"] * (count - 2) + ["closing"]
    return {"title": title, "subtitle": "What it is about",
            "slides": [{"role": roles[i], "title": f"Slide {i + 1}", "gist": f"It shows point {i + 1}."}
                       for i in range(count)]}


def slides_answer(ids, kind="bullets", notes=""):
    """A content answer for the given slide ids."""
    out = []
    for i, sid in enumerate(ids):
        entry = {"id": sid, "title": f"Heading {i + 1}", "subtitle": "A subtitle", "notes": notes,
                 "images": [{"prompt": "a calm landscape"}]}
        if kind == "bullets":
            entry["body"] = {"kind": "bullets",
                             "items": [f"Point {n} of slide {i + 1}" for n in range(1, 7)]}
        elif kind == "table":
            entry["body"] = {"kind": "table", "header": ["Thing", "Value"],
                             "rows": [["One", "1"], ["Two", "2"]]}
        elif kind == "chart":
            entry["body"] = {"kind": "chart", "chart_type": "bar", "categories": ["A", "B"],
                             "series": [{"name": "Count", "values": [1, 2]}]}
        elif kind == "cards":
            entry["body"] = {"kind": "cards", "cards": [
                {"icon": "book-open", "title": "First", "text": "One sentence about it."},
                {"icon": "clock", "title": "Second", "text": "Another sentence about it."}]}
        elif kind == "bars":
            entry["body"] = {"kind": "bars", "bars": [
                {"label": "First", "value": 10, "text": "10 %"},
                {"label": "Second", "value": 40, "text": "40 %"}]}
        elif kind == "figures":
            entry["body"] = {"kind": "figures", "figures": [
                {"value": "8 h", "label": "Per week"}, {"value": "12", "label": "Chapters"}]}
        elif kind == "steps":
            entry["body"] = {"kind": "steps", "steps": [
                {"title": "Plan", "text": "Decide what to do."},
                {"title": "Do", "text": "Carry it out."},
                {"title": "Check", "text": "Look at the result."}]}
        elif kind == "levels":
            entry["body"] = {"kind": "levels", "levels": [
                {"title": "Top", "text": "The highest level."},
                {"title": "Middle", "text": "Below the top."},
                {"title": "Base", "text": "The widest level."}]}
        elif kind == "symbols":
            entry["body"] = {"kind": "symbols", "symbols": [
                {"icon": "book-open", "label": "Reading"}, {"icon": "clock", "label": "Time"},
                {"icon": "flask", "label": "Research"}]}
        elif kind == "highlight":
            entry["body"] = {"kind": "highlight", "key": "42", "text": "The answer is 42 in the end.",
                             "emphasis": "answer"}
        elif kind == "image":
            entry["body"] = {"kind": "image", "prompt": "a quiet library"}
        entry["columns"] = [{"heading": f"Column {c}", "items": [f"Point {c}.{n}" for n in range(1, 5)]}
                            for c in (1, 2)]
        out.append(entry)
    return {"slides": out}


# --- Real documents, built once per test, so the readers run on real files ---------------


def make_pdf(pages: list[list[tuple[int, int, str]]], width: int = 595,
             height: int = 842) -> bytes:
    """A PDF with a text layer, every line placed where it is asked for."""
    font = "/Type/Font/Subtype/Type1/BaseFont/Helvetica/Encoding/WinAnsiEncoding"
    objects: list = ["<</Type/Catalog/Pages 2 0 R>>", "", f"<<{font}>>"]
    kids = []
    for i, lines in enumerate(pages):
        number = 4 + 2 * i
        kids.append(f"{number} 0 R")
        drawn = []
        for x, y, text in lines:
            safe = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            drawn.append(f"BT /F1 11 Tf 1 0 0 1 {x} {y} Tm ({safe}) Tj ET")
        stream = "\n".join(drawn).encode("cp1252")
        objects.append(f"<</Type/Page/Parent 2 0 R/MediaBox[0 0 {width} {height}]"
                       f"/Resources<</Font<</F1 3 0 R>>>>/Contents {number + 1} 0 R>>")
        objects.append(("stream", f"<</Length {len(stream)}>>".encode(), stream))
    objects[1] = f"<</Type/Pages/Kids[{' '.join(kids)}]/Count {len(pages)}>>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode()
        if isinstance(body, tuple):
            out += body[1] + b"\nstream\n" + body[2] + b"\nendstream"
        else:
            out += body.encode("cp1252")
        out += b"\nendobj\n"
    start = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n{start}\n%%EOF\n".encode()
    return bytes(out)


# A handout whose running header carries a page counter and whose footer repeats.
PDF_HEADER = "Leitfaden Waermepumpe | Seite {n} von {count}"
PDF_FOOTER = "Beispielverlag, Ausgabe 2026"


def _handout_pages(count: int = 6) -> list[list[tuple[int, int, str]]]:
    """A handout: a running header with a page counter and a footer on every page."""
    pages = []
    for n in range(1, count + 1):
        pages.append([
            (72, 800, PDF_HEADER.format(n=n, count=count)),
            (72, 740, f"Kapitel {n}: Die Waermequelle"),
            (72, 710, f"Abschnitt {n} erklaert die Jahresarbeitszahl und den Stromverbrauch."),
            (72, 680, f"Die Vorlauftemperatur {n * 5} Grad entscheidet ueber die Effizienz."),
            (72, 60, PDF_FOOTER),
        ])
    return pages


@pytest.fixture
def handout_pdf(tmp_path):
    """A six page PDF with a running header, a page counter and a repeating footer."""
    path = tmp_path / "leitfaden.pdf"
    path.write_bytes(make_pdf(_handout_pages()))
    return path


@pytest.fixture
def handout_docx(tmp_path):
    from docx import Document

    path = tmp_path / "leitfaden.docx"
    doc = Document()
    doc.add_heading("Waermepumpe im Altbau", level=1)
    for n in range(1, 4):
        doc.add_paragraph(f"Abschnitt {n} erklaert die Jahresarbeitszahl und den Heizstab.")
    doc.save(str(path))
    return path


@pytest.fixture
def handout_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "leitfaden.pptx"
    deck = Presentation()
    blank = deck.slide_layouts[6]
    for n in range(1, 7):
        slide = deck.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
        frame = box.text_frame
        frame.text = f"Folie {n}: Die Waermequelle"
        frame.add_paragraph().text = f"Die Jahresarbeitszahl liegt bei {n} Komma fuenf."
        # The same line on every slide, that is what the boilerplate check is for.
        frame.add_paragraph().text = PDF_FOOTER
    deck.save(str(path))
    return path
