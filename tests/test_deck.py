from io import BytesIO

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

from outline2pptx.deck import build_deck, date_text, layout_index
from outline2pptx import colors, library, split, template
from outline2pptx.schemas import (
    Bar, Body, Card, Column, Entry, Figure, Series, Slide, Symbol,
)

TEMPLATES = ["clean", "warm", "forest"]


@pytest.fixture(scope="module")
def clean():
    return library.load("clean")


def deck_of(slides, loaded, **kwargs):
    data, analysis, _ = loaded
    return build_deck(data, analysis, slides, {}, "A Deck", False, **kwargs)


def texts(prs):
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
    return out


def harvest(slide):
    """Every string the slide really carries, plus its pictures and its charts."""
    found = {"text": [], "pictures": 0, "charts": [], "runs": []}
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            found["pictures"] += 1
        if shape.has_text_frame:
            found["text"].append(shape.text_frame.text)
            for paragraph in shape.text_frame.paragraphs:
                found["runs"].extend((r.text, r.font.size, r.font.color) for r in paragraph.runs)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                found["text"].extend(cell.text for cell in row.cells)
        if getattr(shape, "has_chart", False):
            chart = shape.chart
            found["charts"].append({
                "categories": [str(c) for c in chart.plots[0].categories],
                "series": {s.name: list(s.values) for s in chart.series}})
    return found


def strings_of(slide):
    return [t for t in harvest(slide)["text"] if t.strip()]


def reopen(data):
    return Presentation(BytesIO(data))


BODIES = {
    "bullets": Body(kind="bullets", items=[f"Point number {n} on this slide" for n in range(1, 7)]),
    "table": Body(kind="table", header=["Thing", "Value"], rows=[["One", "1"], ["Two", "2"]]),
    "chart": Body(kind="chart", chart_type="bar", categories=["A", "B"],
                  series=[Series(name="Count", values=[1.0, 2.0])]),
    "cards": Body(kind="cards", cards=[Card(icon="book-open", title="Read", text="Start here."),
                                       Card(icon="clock", title="Practice", text="Every day.")]),
    "bars": Body(kind="bars", bars=[Bar(label="First", value=10, text="10 %"),
                                    Bar(label="Second", value=40, text="40 %")]),
    "figures": Body(kind="figures", figures=[Figure(value="8 h", label="Per week"),
                                             Figure(value="12", label="Chapters")]),
    "steps": Body(kind="steps", steps=[Entry(title="Plan", text="Decide what to do."),
                                       Entry(title="Do", text="Carry it out."),
                                       Entry(title="Check", text="Look at it.")]),
    "levels": Body(kind="levels", levels=[Entry(title="Top", text="The highest."),
                                          Entry(title="Middle", text="Below the top."),
                                          Entry(title="Base", text="The widest.")]),
    "symbols": Body(kind="symbols", symbols=[Symbol(icon="book-open", label="Reading"),
                                             Symbol(icon="clock", label="Time"),
                                             Symbol(icon="flask", label="Research")]),
    "highlight": Body(kind="highlight", key="42", text="The answer is 42 in the end.",
                      emphasis="answer"),
}


@pytest.mark.parametrize("name", TEMPLATES)
def test_every_template_has_the_roles_the_planner_needs(name):
    _, analysis, meta = library.load(name)
    assert meta["name"]
    assert meta["style_text"]
    assert sorted(template.selectable_roles(analysis)) == [
        "section", "text", "text_with_image", "two_columns"]
    for role in ("cover", "closing", "text", "section", "two_columns", "text_with_image"):
        assert analysis.roles.get(role), f"{name} has no layout for {role}"
    assert analysis.width == 12192000 and analysis.height == 6858000
    assert analysis.typeface == "Poppins"


@pytest.mark.parametrize("name", TEMPLATES)
def test_every_template_builds_a_deck(name):
    loaded = library.load(name)
    slides = [Slide(id="s1", role="cover", title="A Deck", subtitle="About things"),
              Slide(id="s2", role="section", title="The middle"),
              Slide(id="s3", role="text", title="Points", body=BODIES["bullets"]),
              Slide(id="s4", role="closing", title="One thought", subtitle="Take it home")]
    prs = reopen(deck_of(slides, loaded))
    assert len(prs.slides._sldIdLst) == 4
    assert "A Deck" in texts(prs)
    assert "One thought" in texts(prs)


# The strings of every BODIES fixture that have to come back out of the file.
EXPECTED = {
    "bullets": ["Point number 1 on this slide", "Point number 6 on this slide"],
    "table": ["Thing", "Value", "One", "1", "Two", "2"],
    "chart": [],
    "cards": ["Read", "Start here.", "Practice", "Every day."],
    "bars": ["First", "10 %", "Second", "40 %"],
    "figures": ["8 h", "Per week", "12", "Chapters"],
    "steps": ["Plan", "Decide what to do.", "Do", "Carry it out.", "Check", "Look at it."],
    "levels": ["Top", "The highest.", "Middle", "Below the top.", "Base", "The widest."],
    "symbols": ["Reading", "Time", "Research"],
    "highlight": ["42"],
}
# Icons are pictures, so a card and a symbol slide carries one per entry.
PICTURES = {"cards": 2, "symbols": 3}


@pytest.mark.parametrize("name", TEMPLATES)
@pytest.mark.parametrize("kind", sorted(BODIES))
def test_every_body_kind_puts_its_values_in_the_file(kind, name):
    loaded = library.load(name)
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title=f"A slide of {kind}", body=BODIES[kind]),
              Slide(id="s3", role="closing", title="Done")]
    prs = reopen(deck_of(slides, loaded))
    assert len(prs.slides._sldIdLst) == 3
    found = harvest(prs.slides[1])
    written = " \n".join(found["text"])
    assert f"A slide of {kind}" in written
    for wanted in EXPECTED[kind]:
        assert wanted in written, f"{kind} lost {wanted!r} on {name}"
    assert found["pictures"] == PICTURES.get(kind, 0)


def test_a_chart_carries_its_categories_and_values(clean):
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title="Numbers", body=BODIES["chart"])]
    charts = harvest(reopen(deck_of(slides, clean)).slides[1])["charts"]
    assert len(charts) == 1
    assert charts[0]["categories"] == ["A", "B"]
    assert charts[0]["series"] == {"Count": [1.0, 2.0]}


def test_a_table_header_takes_the_template_colour_and_the_rows_band(clean):
    data, analysis, _ = clean
    palette = colors.from_template(Presentation(BytesIO(data)), analysis)
    body = Body(kind="table", header=["Thing", "Value"],
                rows=[["One", "1"], ["Two", "2"], ["Three", "3"]])
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title="A table", body=body)]
    table = next(sh.table for sh in reopen(deck_of(slides, clean)).slides[1].shapes
                 if getattr(sh, "has_table", False))
    fills = [table.cell(r, 0).fill.fore_color.rgb for r in range(4)]
    assert fills[0] == palette.tint(colors.FULL)
    assert fills[1] == RGBColor(0xFF, 0xFF, 0xFF)
    assert fills[2] == palette.tint(colors.BAND)
    assert fills[3] == RGBColor(0xFF, 0xFF, 0xFF)


def test_a_highlight_marks_its_emphasis_in_the_template_colour(clean):
    data, analysis, _ = clean
    palette = colors.from_template(Presentation(BytesIO(data)), analysis)
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title="One number", body=BODIES["highlight"])]
    runs = [(text, color) for text, _, color in harvest(reopen(deck_of(slides, clean)).slides[1])["runs"]]
    sentence = [r for r in runs if r[0] in ("The ", "answer", " is 42 in the end.")]
    assert [text for text, _ in sentence] == ["The ", "answer", " is 42 in the end."]
    assert sentence[1][1].rgb == palette.text_color()
    assert sentence[0][1].rgb != palette.text_color()


def test_bullets_that_overflow_come_out_smaller(clean):
    _, analysis, _ = clean
    field = analysis.layout(analysis.roles["text"][0]).of_kind("body")[0]
    many = [f"Point number {n}, a sentence long enough on its own that the whole list "
            f"cannot possibly keep the size the template asks for" for n in range(1, 9)]
    few = ["One", "Two"]
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title="Many", body=Body(kind="bullets", items=many)),
              Slide(id="s3", role="text", title="Few", body=Body(kind="bullets", items=few))]
    prs = reopen(deck_of(slides, clean))
    shrunk = [size for _, size, _ in harvest(prs.slides[1])["runs"] if size is not None]
    kept = [size for _, size, _ in harvest(prs.slides[2])["runs"] if size is not None]
    assert shrunk and max(shrunk).pt < field.size_pt
    assert not kept or max(kept).pt >= field.size_pt


def test_a_long_heading_is_shrunk_and_the_body_moves_down(clean):
    _, analysis, _ = clean
    index = analysis.roles["text"][0]
    body_field = analysis.layout(index).of_kind("body")[0]
    long_title = ("A heading so long that it needs three whole lines of the slide before the "
                  "points below it can begin at all")
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", layout=index, title=long_title,
                    body=Body(kind="bullets", items=["One", "Two", "Three"])),
              Slide(id="s3", role="text", layout=index, title="Short",
                    body=Body(kind="bullets", items=["One", "Two", "Three"]))]
    prs = reopen(deck_of(slides, clean))

    def head_and_body(slide):
        boxes = sorted(((sh.top, sh) for sh in slide.shapes if sh.has_text_frame), key=lambda x: x[0])
        return boxes[0][1], boxes[1][1]

    long_head, long_body = head_and_body(prs.slides[1])
    short_head, short_body = head_and_body(prs.slides[2])
    long_size = long_head.text_frame.paragraphs[0].runs[0].font.size
    assert long_size is not None and long_size.pt < analysis.layout(index).of_kind("heading")[0].size_pt
    assert long_body.top > short_body.top
    assert long_body.height < body_field.height


def test_two_columns_writes_both_columns(clean):
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="two_columns", title="Both sides", columns=[
                  Column(heading="Strengths", items=["Fast", "Cheap", "Proven"]),
                  Column(heading="Limits", items=["Needs data", "Can overfit", "Opaque"])]),
              Slide(id="s3", role="closing", title="Done")]
    written = texts(reopen(deck_of(slides, clean)))
    assert any("Cheap" in t for t in written)
    assert any("Can overfit" in t for t in written)


def test_a_layout_with_subheadings_keeps_them(clean):
    _, analysis, _ = clean
    with_heads = next(i for i in analysis.roles["two_columns"]
                      if analysis.layout(i).of_kind("subheading"))
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="two_columns", layout=with_heads, title="Both sides", columns=[
                  Column(heading="Strengths", items=["Fast", "Cheap", "Proven"]),
                  Column(heading="Limits", items=["Needs data", "Can overfit", "Opaque"])]),
              Slide(id="s3", role="closing", title="Done")]
    written = texts(reopen(deck_of(slides, clean)))
    assert any("Strengths" in t for t in written)
    assert any("Limits" in t for t in written)


def test_notes_only_reach_the_file_when_asked(clean):
    data, analysis, _ = clean
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title="Points", body=BODIES["bullets"], notes="Say this."),
              Slide(id="s3", role="closing", title="Done")]
    with_notes = reopen(build_deck(data, analysis, slides, {}, "A Deck", True))
    assert "Say this." in with_notes.slides[1].notes_slide.notes_text_frame.text
    without = reopen(build_deck(data, analysis, slides, {}, "A Deck", False))
    assert not without.slides[1].has_notes_slide


def test_the_notice_lands_on_every_slide(clean):
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="closing", title="Done")]
    written = texts(reopen(deck_of(slides, clean, notice="Built with a model")))
    assert sum(1 for t in written if t == "Built with a model") == 2


def boxes_of(prs):
    """Name, top, height and text of every text shape of every slide."""
    return [(shape.name, shape.top, shape.height, shape.text_frame.text)
            for slide in prs.slides for shape in slide.shapes if shape.has_text_frame]


@pytest.mark.parametrize("name", TEMPLATES)
def test_the_bottom_strip_carries_date_number_and_presenter(name):
    loaded = library.load(name)
    slides = [Slide(id="s1", role="cover", title="A Deck", subtitle="About things"),
              Slide(id="s2", role="text", title="Points", body=BODIES["bullets"])]
    prs = reopen(deck_of(slides, loaded, date="Mar 15, 2026", presenter="Jane"))
    found = boxes_of(prs)
    assert [t for *_, t in found if t == "Jane"], "the presenter never reached the cover"
    assert [t for *_, t in found if t == "Mar 15, 2026"], "the date never reached the slide"
    assert [t for *_, t in found if t == "2"], "the slide number never reached the slide"
    for shape_name, top, height, _ in found:
        assert top and height, f"{shape_name} has no box of its own"


@pytest.mark.parametrize("name", TEMPLATES)
def test_the_bottom_strip_sits_below_the_body_area(name):
    data, analysis, _ = library.load(name)
    prs = reopen(data)
    strip = {"DATE (16)", "FOOTER (15)", "SLIDE_NUMBER (13)"}
    for layout in prs.slide_masters[0].slide_layouts:
        boxes = {str(ph.placeholder_format.type): ph for ph in layout.placeholders}
        assert strip <= set(boxes), f"{layout.name} misses part of the bottom strip"
        content = [ph for key, ph in boxes.items() if key not in strip]
        lowest = max((ph.top + ph.height for ph in content), default=0)
        for key in strip:
            ph = boxes[key]
            assert ph.top and ph.height, f"{key} of {layout.name} has no box of its own"
            assert ph.top >= lowest, f"{key} of {layout.name} runs into the content"
            assert ph.top + ph.height <= analysis.height


def test_the_presenter_only_lands_on_the_cover(clean):
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title="Points", body=BODIES["bullets"])]
    prs = reopen(deck_of(slides, clean, presenter="Jane"))
    first = [t for *_, t in boxes_of(prs) if t == "Jane"]
    assert len(first) == 1
    assert "Jane" in [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]


def test_date_text_follows_the_language():
    from datetime import date

    assert date_text("en", date(2026, 3, 15)) == "Mar 15, 2026"
    assert date_text("de", date(2026, 3, 15)) == "15.03.2026"


def test_layout_index_is_kept_once_chosen(clean):
    _, analysis, _ = clean
    s = Slide(id="s1", role="text", layout=analysis.roles["text"][-1])
    assert layout_index(analysis, s, {}) == analysis.roles["text"][-1]


def test_layout_rotates_between_the_layouts_of_a_role(clean):
    _, analysis, _ = clean
    counter: dict = {}
    picked = [layout_index(analysis, Slide(id=f"s{i}", role="text"), counter) for i in range(4)]
    assert len(set(picked)) == len(analysis.roles["text"])


# Splitting


def test_overflowing_columns_are_split(clean):
    _, analysis, _ = clean
    long = [f"This is point number {n} and it carries a full sentence of its own, twice over, "
            f"so that it certainly wraps" for n in range(1, 9)]
    s = Slide(id="s1", role="two_columns", title="Many points",
              columns=[Column(heading="Left", items=list(long)),
                       Column(heading="Right", items=list(long))],
              notes="First part. Second part. Third part. Fourth part.")
    out = split.split_overflow([s], analysis)
    assert len(out) > 1
    assert sum(len(c.items) for p in out for c in p.columns) == 2 * len(long)
    assert all(p.notes for p in out)
    assert len({p.id for p in out}) == len(out)


def test_a_slide_that_fits_is_not_split(clean):
    _, analysis, _ = clean
    s = Slide(id="s1", role="text", title="Few", body=Body(kind="bullets", items=["One", "Two"]))
    assert split.split_overflow([s], analysis) == [s]


def test_new_id_counts_past_the_highest():
    assert split.new_id({"s1", "s7", "s3"}) == "s8"


def test_notes_are_split_at_a_sentence_end():
    first, second = split.split_notes("One sentence here. Another sentence here.", 0.5)
    assert first == "One sentence here."
    assert second == "Another sentence here."


# Bullets


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_MARKERS = ("buChar", "buAutoNum", "buBlip")


def _bullet_of(level):
    """'none', a marker name, or None when this level says nothing about bullets."""
    if level is None:
        return None
    if level.find(f"{_A}buNone") is not None:
        return "none"
    for name in _MARKERS:
        if level.find(f"{_A}{name}") is not None:
            return name
    return None


def shows_a_bullet(shape) -> bool:
    """What the first paragraph of this placeholder really shows: its own rule, the layout's, the master's."""
    own = _bullet_of(shape.text_frame.paragraphs[0]._p.find(f"{_A}pPr"))
    if own is None:
        idx = shape.placeholder_format.idx
        layout_ph = next((p for p in shape.part.slide_layout.placeholders
                          if p.placeholder_format.idx == idx), None)
        if layout_ph is not None:
            own = _bullet_of(layout_ph._element.find(f".//{_A}lstStyle/{_A}lvl1pPr"))
    if own is None:
        master = shape.part.slide_layout.slide_master._element
        own = _bullet_of(master.find(f".//{_P}bodyStyle/{_A}lvl1pPr"))
    return own not in (None, "none")


@pytest.mark.parametrize("name", TEMPLATES)
def test_one_line_placeholders_switch_the_bullet_off(name):
    data, analysis, _ = library.load(name)
    prs = reopen(data)
    for layout in analysis.layouts:
        pptx_layout = prs.slide_masters[0].slide_layouts[layout.index]
        for one_liner in layout.of_kind("heading") + layout.of_kind("subheading"):
            ph = next(p for p in pptx_layout.placeholders
                      if p.placeholder_format.idx == one_liner.idx)
            level = ph._element.find(f".//{_A}lstStyle/{_A}lvl1pPr")
            assert _bullet_of(level) == "none", (
                f"{name}: {layout.name} placeholder {one_liner.idx} keeps the master's bullet")


@pytest.mark.parametrize("name", TEMPLATES)
def test_the_cover_line_has_no_bullet_but_a_list_does(name):
    loaded = library.load(name)
    slides = [Slide(id="s1", role="cover", title="A Deck", subtitle="About things"),
              Slide(id="s2", role="text", title="Points", body=BODIES["bullets"]),
              Slide(id="s3", role="closing", title="One thought", subtitle="Take it home")]
    prs = reopen(deck_of(slides, loaded))
    cover = next(sh for sh in prs.slides[0].shapes if sh.text_frame.text == "About things")
    closing = next(sh for sh in prs.slides[2].shapes if sh.text_frame.text == "Take it home")
    points = next(sh for sh in prs.slides[1].shapes
                  if "Point number 1 on this slide" in sh.text_frame.text)
    assert not shows_a_bullet(cover)
    assert not shows_a_bullet(closing)
    assert shows_a_bullet(points)
