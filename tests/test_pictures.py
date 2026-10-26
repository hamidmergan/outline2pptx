import shutil
from io import BytesIO

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from outline2pptx import config, image, library, pictures, preview, render
from outline2pptx.deck import build_deck
from outline2pptx.errors import PreviewError
from outline2pptx.schemas import Body, Column, ImageSpec, Slide

HAS_LIBREOFFICE = render.available()


@pytest.fixture(scope="module")
def clean():
    return library.load("clean")


def settings_for(server=None, images=True):
    if server is None:
        return config.resolve(base_url="http://x/v1", model="m", images=False)
    return config.resolve(base_url=server.base_url, model="m",
                          image_base_url=server.base_url, image_model="draw", images=images)


def picture_slides(analysis):
    index = analysis.roles["text_with_image"][0]
    return [Slide(id="s1", role="text_with_image", layout=index, title="A picture slide",
                  images=[ImageSpec(prompt="a quiet library")])], {"s1": index}


def test_image_size_follows_the_field():
    assert image.image_size(5303520, 4300000) == "1744x1408"
    # a tiny field still gets a usable picture, a huge one stays capped
    assert image.image_size(10000, 10000) == "768x768"
    assert image.image_size(12192000, 1000000) == "1920x256"


def test_a_quoted_word_is_taken_out_of_a_prompt():
    assert '"Quiet"' not in image.prompt_for("Flat.", 'A door with a sign "Quiet" on it')
    assert image.prompt_for("Flat.", "A door").startswith("Flat. Subject: A door")
    assert image.NO_LETTERING in image.prompt_for("", "A door")


def test_without_an_image_model_the_slide_keeps_no_picture_box(clean):
    """A deck built without pictures must not leave a coloured slab to delete by hand."""
    data, analysis, _ = clean
    slides, layouts = picture_slides(analysis)
    drawn = pictures.draw(analysis, slides, layouts, settings_for(), "Flat.", "2563EB")
    assert drawn == {}
    prs = Presentation(BytesIO(build_deck(data, analysis, slides, drawn, "A Deck", False)))
    shapes = list(prs.slides[0].shapes)
    assert "A picture slide" in [sh.text_frame.text for sh in shapes if sh.has_text_frame]
    assert not any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in shapes)
    assert not any(sh.is_placeholder and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE
                   for sh in shapes)


def test_the_image_model_draws_every_field(clean, server):
    _, analysis, _ = clean
    slides, layouts = picture_slides(analysis)
    drawn = pictures.draw(analysis, slides, layouts, settings_for(server), "Flat.", "2563EB")
    assert set(drawn) == {"s1_0"}
    assert drawn["s1_0"] == server.image_bytes
    assert server.kinds() == ["image"]
    body = server.requests[0]["body"]
    assert body["model"] == "draw"
    assert body["size"] == image.image_size(*_field_size(analysis))
    assert "a quiet library" in body["prompt"]


def _field_size(analysis):
    layout = analysis.layout(analysis.roles["text_with_image"][0])
    field = layout.of_kind("picture")[0]
    return field.width, field.height


def test_a_refused_size_is_retried_as_a_square(clean, server):
    _, analysis, _ = clean
    slides, layouts = picture_slides(analysis)
    server.refuse_images = 1
    drawn = pictures.draw(analysis, slides, layouts, settings_for(server), "Flat.", "2563EB")
    assert [r["body"]["size"] for r in server.requests] == [
        image.image_size(*_field_size(analysis)), "1024x1024"]
    # The square was answered, so the drawn picture is the model's, not a placeholder.
    assert drawn["s1_0"] == server.image_bytes
    assert drawn["s1_0"] != image.placeholder(*_field_size(analysis), "2563EB")


def test_both_sizes_refused_leaves_a_placeholder(clean, server):
    _, analysis, _ = clean
    slides, layouts = picture_slides(analysis)
    server.image_status = 400
    drawn = pictures.draw(analysis, slides, layouts, settings_for(server), "Flat.", "2563EB")
    assert len(server.requests) == 2
    assert drawn["s1_0"] != server.image_bytes
    assert drawn["s1_0"] == image.placeholder(*_field_size(analysis), "2563EB")


def test_a_failing_image_model_falls_back_to_a_placeholder(clean, server):
    _, analysis, _ = clean
    slides, layouts = picture_slides(analysis)
    server.fail = True
    drawn = pictures.draw(analysis, slides, layouts, settings_for(server), "Flat.", "2563EB")
    assert drawn["s1_0"].startswith(b"\x89PNG")


def test_a_drawn_picture_reaches_the_deck(clean, server):
    data, analysis, _ = clean
    slides, layouts = picture_slides(analysis)
    drawn = pictures.draw(analysis, slides, layouts, settings_for(server), "Flat.", "2563EB")
    deck = build_deck(data, analysis, slides, drawn, "A Deck", False)
    prs = Presentation(BytesIO(deck))
    pictures_on_slide = [s for s in prs.slides[0].shapes if hasattr(s, "image")]
    assert len(pictures_on_slide) == 1
    assert pictures_on_slide[0].image.blob


def test_images_need_a_model_in_the_settings():
    from outline2pptx.errors import ConfigError

    with pytest.raises(ConfigError, match="no image model"):
        config.resolve(base_url="http://x/v1", model="m", images=True)


@pytest.mark.skipif(HAS_LIBREOFFICE, reason="LibreOffice is installed here")
def test_preview_says_so_when_libreoffice_is_missing(clean):
    data, analysis, _ = clean
    deck = build_deck(data, analysis, [Slide(id="s1", role="cover", title="A")], {}, "A", False)
    with pytest.raises(PreviewError, match="LibreOffice"):
        preview.render(deck)


@pytest.mark.skipif(not HAS_LIBREOFFICE, reason="LibreOffice is not installed here")
def test_preview_renders_one_png_per_slide(clean):
    data, analysis, _ = clean
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="closing", title="Done")]
    deck = build_deck(data, analysis, slides, {}, "A Deck", False)
    pngs = preview.render(deck)
    assert len(pngs) == 2
    assert all(p.startswith(b"\x89PNG") for p in pngs)


def test_available_checks_the_three_tools(monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/" + name)
    assert render.available()
    monkeypatch.setattr(shutil, "which", lambda name: None if name == "pdftoppm" else "/usr/bin/x")
    assert not render.available()


def _picture_only_analysis():
    """A template whose only body layout carries a picture, so it has no plain text role."""
    from outline2pptx.template import Analysis, Field, Layout

    def head(idx=0):
        return Field(idx=idx, kind="heading", ph_type="TITLE", left=640080, top=457200,
                     width=10911840, height=1000000, size_pt=32.0)

    cover = Layout(index=0, name="Cover", role="cover", fields=[head()])
    picture = Layout(index=1, name="Picture", role="text_with_image", fields=[
        head(),
        Field(idx=1, kind="body", ph_type="OBJECT", left=640080, top=1700000, width=5303520,
              height=4300000, size_pt=18.0, is_object=True),
        Field(idx=2, kind="picture", ph_type="PICTURE", left=6248400, top=1700000,
              width=5303520, height=4300000)])
    return Analysis(width=12192000, height=6858000, typeface="Poppins",
                    layouts=[cover, picture],
                    roles={"cover": [0], "closing": [0], "text_with_image": [1]})


def test_a_split_picture_slide_still_gets_its_pictures():
    from outline2pptx import split

    analysis = _picture_only_analysis()
    long = [f"Point number {n}, a whole sentence of its own so that the column cannot hold it"
            for n in range(1, 9)]
    s = Slide(id="s1", role="text_with_image", layout=1, title="A picture slide",
              columns=[Column(heading="Left", items=list(long))],
              images=[ImageSpec(prompt="a quiet library")])
    ordered = split.split_overflow([s], analysis)
    assert len(ordered) == 2
    # The rest has no text layout to fall back to, but it still carries one.
    assert all(part.layout is not None for part in ordered)
    layouts = {part.id: part.layout for part in ordered if part.layout is not None}
    assert [key for key, _, _ in pictures.plan(analysis, ordered, layouts)] == ["s1_0"]


def test_a_slide_without_a_stored_layout_rotates_instead_of_failing(clean):
    _, analysis, _ = clean
    slides, _ = picture_slides(analysis)
    slides[0].layout = None
    assert [key for key, _, _ in pictures.plan(analysis, slides, {})] == ["s1_0"]


def test_no_slide_image_is_a_preview_error_not_a_raw_runtime_error(clean, tmp_path, monkeypatch):
    """LibreOffice wrote its PDF but pdftoppm produced nothing, a full disk for instance."""
    data, analysis, _ = clean
    deck = build_deck(data, analysis, [Slide(id="s1", role="cover", title="A")], {}, "A", False)
    source = tmp_path / "deck.pptx"
    source.write_bytes(deck)
    target = tmp_path / "img"
    target.mkdir()

    class Done:
        returncode = 0
        stdout = "Pages: 1\n"
        stderr = ""

    def fake_run(cmd, **kwargs):
        if "pdftoppm" not in cmd[0] and "pdfinfo" not in cmd[0]:
            (target / "_render_deck.pdf").write_bytes(b"%PDF-1.4\n")
        return Done()

    monkeypatch.setattr(render.subprocess, "run", fake_run)
    with pytest.raises(PreviewError, match="no slide image"):
        render.pptx_to_images(source, target)


def test_the_preview_finds_libreoffice_under_its_macos_name(monkeypatch):
    from outline2pptx import render

    found = {"soffice", "pdfinfo", "pdftoppm"}
    monkeypatch.setattr(render.shutil, "which", lambda name: f"/bin/{name}" if name in found else None)
    assert render.available()
    assert render.build_soffice_cmd(render.Path("deck.pptx"), render.Path("out"))[0] == "soffice"
