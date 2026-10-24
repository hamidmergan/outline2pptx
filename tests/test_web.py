import json
import time
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from conftest import PDF_FOOTER, outline_answer
from outline2pptx import api
from test_api import OUTLINE_MD, answer_for_batch
from test_material_outline import echo_slides, material_lines

pytest.importorskip("fastapi")


def make_client(server=None, **extra):
    from outline2pptx.web.app import create_app

    defaults = dict(extra)
    if server is not None:
        defaults.setdefault("base_url", server.base_url)
        defaults.setdefault("model", "a-model")
    return TestClient(create_app(defaults, allowed_hosts=["testserver"]))


def wait_for(client, state="done", seconds=30):
    deadline = time.time() + seconds
    while time.time() < deadline:
        job = client.get("/api/status").json()
        if job["state"] == state:
            return job
        if job["state"] == "error":
            return job
        time.sleep(0.05)
    raise AssertionError(f"the job never reached {state}")


def test_the_page_is_served():
    client = make_client()
    response = client.get("/")
    assert response.status_code == 200
    assert "outline2pptx" in response.text
    assert client.get("/static/app.js").status_code == 200
    assert client.get("/static/style.css").status_code == 200


def test_info_lists_the_templates_and_the_limits():
    data = make_client().get("/api/info").json()
    assert [t["id"] for t in data["templates"]] == ["clean", "forest", "warm"]
    assert data["languages"] == ["de", "en"]
    assert data["min_slides"] == 3 and data["max_slides"] == 30
    assert data["configured"] is False
    # What the documents field on the page needs to know.
    assert ".pdf" in data["extensions"] and ".docx" in data["extensions"]
    assert data["max_material_files"] == 10
    assert data["max_material_chars"] == 500_000
    assert data["window_chars"] > 0


def test_info_shows_what_the_environment_decided(server):
    data = make_client(server, language="de", template="forest").get("/api/info").json()
    assert data["configured"] is True
    assert data["model"] == "a-model"
    assert data["language"] == "de"
    assert data["template"] == "forest"


def test_an_unknown_host_is_refused():
    from outline2pptx.web.app import create_app

    client = TestClient(create_app({}, allowed_hosts=["localhost"]))
    assert client.get("/", headers={"Host": "evil.example"}).status_code == 400


def test_a_cross_origin_request_is_refused():
    client = make_client()
    response = client.get("/api/info", headers={"Origin": "http://evil.example"})
    assert response.status_code == 403


def test_the_outline_route_plans_from_a_topic(server):
    server.outline = outline_answer(6, title="Bees")
    client = make_client(server)
    data = client.post("/api/outline", json={"topic": "Bees", "slides": 6}).json()
    assert data["outline"].startswith("# Bees")
    assert server.kinds() == ["outline"]


def test_the_outline_route_needs_a_topic():
    response = make_client().post("/api/outline", json={"topic": "  "})
    assert response.status_code == 400
    assert "topic" in response.json()["detail"]


def test_the_fields_on_the_page_beat_the_environment(server):
    server.outline = outline_answer(5)
    client = make_client(base_url="http://127.0.0.1:1/v1", model="wrong")
    client.post("/api/outline", json={"topic": "Bees", "base_url": server.base_url,
                                      "model": "from-the-page"})
    assert server.requests[0]["body"]["model"] == "from-the-page"


def test_without_a_model_the_outline_route_says_what_is_missing():
    response = make_client().post("/api/outline", json={"topic": "Bees"})
    assert response.status_code == 400
    assert "base_url" in response.json()["detail"]


def test_build_runs_and_the_deck_can_be_downloaded(server):
    server.slides = answer_for_batch
    client = make_client(server)
    assert client.post("/api/build", json={"outline": OUTLINE_MD}).json()["state"] == "running"
    job = wait_for(client)
    assert job["state"] == "done", job["error"]
    assert len(job["slides"]) == 6
    assert job["slides"][0]["number"] == 1 and job["slides"][0]["role"] == "cover"
    assert job["info"]["template"] == "Clean"
    deck = client.get("/api/download")
    assert deck.status_code == 200
    assert deck.headers["content-disposition"].endswith('filename="deck.pptx"')
    prs = Presentation(BytesIO(deck.content))
    assert len(prs.slides._sldIdLst) == 6


def test_an_empty_outline_is_refused(server):
    response = make_client(server).post("/api/build", json={"outline": "  "})
    assert response.status_code == 400


def test_a_failing_model_ends_the_job_with_the_reason(server):
    server.slides = "nonsense"
    client = make_client(server)
    client.post("/api/build", json={"outline": OUTLINE_MD})
    job = wait_for(client, "error")
    assert job["state"] == "error" and job["error"]


def test_images_need_an_image_model(server):
    response = make_client(server).post("/api/build", json={"outline": OUTLINE_MD, "images": True})
    assert response.status_code == 400
    assert "image_model" in response.json()["detail"]


def test_notes_reach_the_file(server):
    from conftest import slides_answer

    def reply(body):
        user = body["messages"][1]["content"]
        ids = [line.split()[1] for line in user.splitlines() if line.startswith("SLIDE ")]
        return slides_answer(ids, notes="What the speaker says here.")

    server.slides = reply
    client = make_client(server)
    client.post("/api/build", json={"outline": OUTLINE_MD, "notes": True})
    assert wait_for(client)["state"] == "done"
    prs = Presentation(BytesIO(client.get("/api/download").content))
    assert any(s.has_notes_slide for s in prs.slides)


def test_download_before_a_run_is_a_404():
    assert make_client().get("/api/download").status_code == 404


def test_the_preview_says_so_without_libreoffice(server, monkeypatch):
    server.slides = answer_for_batch
    monkeypatch.setattr("outline2pptx.render.available", lambda: False)
    client = make_client(server)
    client.post("/api/build", json={"outline": OUTLINE_MD})
    assert wait_for(client)["state"] == "done"
    response = client.get("/api/preview/1")
    assert response.status_code == 409
    assert "LibreOffice" in response.json()["detail"]


def test_the_preview_serves_one_png_per_slide(server, monkeypatch):
    server.slides = answer_for_batch
    monkeypatch.setattr("outline2pptx.render.available", lambda: True)
    monkeypatch.setattr("outline2pptx.preview.render", lambda data, dpi=80: [b"\x89PNG-one"])
    client = make_client(server)
    client.post("/api/build", json={"outline": OUTLINE_MD})
    assert wait_for(client)["state"] == "done"
    response = client.get("/api/preview/1")
    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    assert client.get("/api/preview/9").status_code == 404


def test_a_second_outline_is_refused_while_one_is_being_planned(server, monkeypatch):
    """One call to the model at a time, so two browser tabs cannot plan two outlines at once."""
    import threading

    started, go = threading.Event(), threading.Event()

    def slow(topic, **kwargs):
        started.set()
        assert go.wait(10)
        return OUTLINE_MD

    monkeypatch.setattr("outline2pptx.api.outline_from_topic", slow)
    client = make_client(server)
    first: dict = {}

    def plan():
        first["response"] = client.post("/api/outline", json={"topic": "Bees"})

    worker = threading.Thread(target=plan)
    worker.start()
    try:
        assert started.wait(10)
        second = client.post("/api/outline", json={"topic": "Bees"})
        build = client.post("/api/build", json={"outline": OUTLINE_MD})
    finally:
        go.set()
        worker.join(10)
    assert second.status_code == 409
    assert build.status_code == 409
    assert first["response"].status_code == 200
    # The flag is cleared again, so the next call goes through.
    assert client.post("/api/outline", json={"topic": "Bees"}).status_code == 200


def test_the_page_shows_every_field_of_the_deck_info(server):
    """The info the server computes is shown, not collected and dropped."""
    from pathlib import Path

    import outline2pptx.web.app as web

    server.slides = answer_for_batch
    client = make_client(server, instructions="Keep it short", style_text="Flat shapes")
    client.post("/api/build", json={"outline": OUTLINE_MD})
    info = wait_for(client)["info"]
    assert info["model"] == "a-model"
    assert info["template"] and info["language"] and info["instructions"] and info["style_text"]
    page = (Path(web.STATIC) / "index.html").read_text("utf8")
    script = (Path(web.STATIC) / "app.js").read_text("utf8")
    assert 'id="deckInfo"' in page
    for key in info:
        if key == "slides":
            continue
        assert f'"{key}"' in script, f"the page never reads {key}"


# --- Documents on the page ---------------------------------------------------------------

def upload(client, paths, **form):
    files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in paths]
    data = {"fields": json.dumps(form.pop("fields", {})), **{k: str(v) for k, v in form.items()}}
    return client.post("/api/material", files=files, data=data)


def test_the_page_offers_all_three_inputs():
    page = make_client().get("/").text
    assert 'data-tab="topic"' in page
    assert 'data-tab="material"' in page
    assert 'data-tab="markdown"' in page
    assert 'id="files"' in page and "multiple" in page
    # And it says which one is in use.
    assert 'id="inUse"' in page
    script = make_client().get("/static/app.js").text
    assert "inUseMaterial" in script and "materialRead" in script


def test_uploaded_documents_are_read_and_planned_into_an_outline(server, handout_pdf):
    server.outline = material_lines
    client = make_client(server)
    response = upload(client, [handout_pdf])
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["outline"].startswith("# ")
    # The size of what was read reaches the page, together with the window it is measured against.
    assert data["material"]["files"] == [{"name": "leitfaden.pdf",
                                         "chars": data["material"]["files"][0]["chars"],
                                         "pages": 6}]
    assert data["material"]["chars"] > 0
    assert data["material"]["pages"] == 6
    assert data["shortened"] is False
    assert data["window_chars"] > data["material"]["chars"]
    # The handout reached the model, its repeating footer did not.
    sent = server.requests[0]["body"]["messages"][1]["content"]
    assert "Jahresarbeitszahl" in sent
    assert PDF_FOOTER not in sent


def test_the_page_is_told_when_the_material_was_planned_through_its_topics(server, tmp_path):
    path = tmp_path / "lang.txt"
    path.write_text("\n\n".join(f"Abschnitt {n}: Die Jahresarbeitszahl steigt. "
                                 + "Fuellsatz. " * 200 for n in range(1, 40)))
    server.topics = lambda body: json.dumps({"topics": [
        {"title": "Quelle", "gist": "Was das Material sagt."},
        {"title": "Zahl", "gist": "Welche Zahlen es nennt."},
        {"title": "Heizstab", "gist": "Wann er anspringt."}]})
    server.outline = material_lines
    client = make_client(server, context_window=32_768)
    data = upload(client, [path]).json()
    assert data["shortened"] is True
    assert data["material"]["chars"] > data["window_chars"] > 0
    assert "topics" in server.kinds()


def test_several_documents_can_be_uploaded_at_once(server, handout_pdf, handout_docx):
    server.outline = material_lines
    data = upload(make_client(server), [handout_pdf, handout_docx]).json()
    assert [f["name"] for f in data["material"]["files"]] == ["leitfaden.pdf", "leitfaden.docx"]
    system = server.requests[0]["body"]["messages"][0]["content"]
    assert "comes from 2 files" in system


def test_a_number_of_slides_from_the_page_reaches_the_call(server, handout_docx):
    server.outline = material_lines
    upload(make_client(server), [handout_docx], slides=15)
    assert "Plan around 15 slides" in server.requests[0]["body"]["messages"][0]["content"]


def test_a_document_the_deck_cannot_use_is_refused(server, tmp_path):
    path = tmp_path / "sound.mp3"
    path.write_bytes(b"not a document")
    response = upload(make_client(server), [path])
    assert response.status_code == 400
    assert ".mp3" in response.json()["detail"]
    assert server.requests == []


def test_material_over_the_limit_is_refused_with_the_setting(server, tmp_path):
    path = tmp_path / "lang.txt"
    path.write_text("Ein Satz ueber die Waermepumpe. " * 400)
    client = make_client(server, max_material_chars=1000)
    response = upload(client, [path])
    assert response.status_code == 400
    assert "OUTLINE2PPTX_MAX_MATERIAL_CHARS" in response.json()["detail"]


def test_an_upload_without_a_model_says_what_is_missing(handout_docx):
    response = upload(make_client(), [handout_docx])
    assert response.status_code == 400
    assert "base_url" in response.json()["detail"]


def test_the_deck_from_uploaded_documents_can_be_downloaded(server, handout_pdf):
    server.outline = material_lines
    server.slides = echo_slides
    client = make_client(server)
    outline_md = upload(client, [handout_pdf]).json()["outline"]
    assert client.post("/api/build", json={"outline": outline_md}).json()["state"] == "running"
    job = wait_for(client)
    assert job["state"] == "done", job["error"]
    deck = client.get("/api/download")
    prs = Presentation(BytesIO(deck.content))
    texts = [shape.text_frame.text for slide in prs.slides for shape in slide.shapes
             if shape.has_text_frame]
    assert any("Waermequelle" in text for text in texts)
    assert all(PDF_FOOTER not in text for text in texts)


def test_an_upload_while_a_build_runs_is_refused(server, handout_docx, monkeypatch):
    """One run at a time, so a second tab cannot upload while the slides are being written."""
    import threading

    started, go = threading.Event(), threading.Event()
    real = api.build

    def slow(text, *args, **kwargs):
        started.set()
        assert go.wait(10)
        return real(text, *args, **kwargs)

    server.slides = answer_for_batch
    monkeypatch.setattr("outline2pptx.api.build", slow)
    client = make_client(server)
    client.post("/api/build", json={"outline": OUTLINE_MD})
    assert started.wait(10)
    response = upload(client, [handout_docx])
    go.set()
    assert response.status_code == 409
    assert "right now" in response.json()["detail"]
    assert wait_for(client)["state"] == "done"
