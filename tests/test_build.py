from io import BytesIO

import pytest
from pptx import Presentation

from outline2pptx import build, library, split, template
from outline2pptx.schemas import (
    Bar, Body, Card, Column, Entry, Figure, Series, Slide, Symbol,
)

TEMPLATES = ["clean", "warm", "forest"]


@pytest.fixture(scope="module")
def clean():
    return library.load("clean")


def deck_of(slides, loaded, **kwargs):
    data, analysis, _ = loaded
    return build.build_deck(data, analysis, slides, {}, "A Deck", False, **kwargs)


def texts(prs):
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
    return out


def reopen(data):
    return Presentation(BytesIO(data))


BODIES = {
    "bullets": Body(kind="bullets", items=[f"Point number {n} on this slide" for n in range(1, 7)]),
    "table": Body(kind="table", header=["Thing", "Value"], rows=[["One", "1"], ["Two", "2"]]),
    "chart": Body(kind="chart", chart_type="bar", categories=["A", "B"],
                  series=[Series(name="Count", values=[1.0, 2.0])]),
    "cards": Body(kind="cards", cards=[Card(icon="book-open", title="Read", text="Start here."),
                                       Card(icon="clock", title="Practice", text="Every day.")]),
    "bars": Body(kind="bars", bars=[Bar(label="First", value=10, text="10 %"),
                                    Bar(label="Second", value=40, text="40 %")]),
    "figures": Body(kind="figures", figures=[Figure(value="8 h", label="Per week"),
                                             Figure(value="12", label="Chapters")]),
    "steps": Body(kind="steps", steps=[Entry(title="Plan", text="Decide what to do."),
                                       Entry(title="Do", text="Carry it out."),
                                       Entry(title="Check", text="Look at it.")]),
    "levels": Body(kind="levels", levels=[Entry(title="Top", text="The highest."),
                                          Entry(title="Middle", text="Below the top."),
                                          Entry(title="Base", text="The widest.")]),
    "symbols": Body(kind="symbols", symbols=[Symbol(icon="book-open", label="Reading"),
                                             Symbol(icon="clock", label="Time"),
                                             Symbol(icon="flask", label="Research")]),
    "highlight": Body(kind="highlight", key="42", text="The answer is 42 in the end.",
                      emphasis="answer"),
}


@pytest.mark.parametrize("name", TEMPLATES)
def test_every_template_has_the_roles_the_planner_needs(name):
    _, analysis, meta = library.load(name)
    assert meta["name"]
    assert meta["style_text"]
    assert sorted(template.selectable_roles(analysis)) == [
        "section", "text", "text_with_image", "two_columns"]
    for role in ("cover", "closing", "text", "section", "two_columns", "text_with_image"):
        assert analysis.roles.get(role), f"{name} has no layout for {role}"
    assert analysis.width == 12192000 and analysis.height == 6858000
    assert analysis.typeface == "Poppins"


@pytest.mark.parametrize("name", TEMPLATES)
def test_every_template_builds_a_deck(name):
    loaded = library.load(name)
    slides = [Slide(id="s1", role="cover", title="A Deck", subtitle="About things"),
              Slide(id="s2", role="section", title="The middle"),
              Slide(id="s3", role="text", title="Points", body=BODIES["bullets"]),
              Slide(id="s4", role="closing", title="One thought", subtitle="Take it home")]
    prs = reopen(deck_of(slides, loaded))
    assert len(prs.slides._sldIdLst) == 4
    assert "A Deck" in texts(prs)
    assert "One thought" in texts(prs)


@pytest.mark.parametrize("kind", sorted(BODIES))
def test_every_body_kind_builds(kind, clean):
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title=f"A slide of {kind}", body=BODIES[kind]),
              Slide(id="s3", role="closing", title="Done")]
    prs = reopen(deck_of(slides, clean))
    assert len(prs.slides._sldIdLst) == 3
    shapes = list(prs.slides[1].shapes)
    assert len(shapes) > 1


def test_two_columns_writes_both_columns(clean):
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="two_columns", title="Both sides", columns=[
                  Column(heading="Strengths", items=["Fast", "Cheap", "Proven"]),
                  Column(heading="Limits", items=["Needs data", "Can overfit", "Opaque"])]),
              Slide(id="s3", role="closing", title="Done")]
    written = texts(reopen(deck_of(slides, clean)))
    assert any("Cheap" in t for t in written)
    assert any("Can overfit" in t for t in written)


def test_a_layout_with_subheadings_keeps_them(clean):
    _, analysis, _ = clean
    with_heads = next(i for i in analysis.roles["two_columns"]
                      if analysis.layout(i).of_kind("subheading"))
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="two_columns", layout=with_heads, title="Both sides", columns=[
                  Column(heading="Strengths", items=["Fast", "Cheap", "Proven"]),
                  Column(heading="Limits", items=["Needs data", "Can overfit", "Opaque"])]),
              Slide(id="s3", role="closing", title="Done")]
    written = texts(reopen(deck_of(slides, clean)))
    assert any("Strengths" in t for t in written)
    assert any("Limits" in t for t in written)


def test_notes_only_reach_the_file_when_asked(clean):
    data, analysis, _ = clean
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="text", title="Points", body=BODIES["bullets"], notes="Say this."),
              Slide(id="s3", role="closing", title="Done")]
    with_notes = reopen(build.build_deck(data, analysis, slides, {}, "A Deck", True))
    assert "Say this." in with_notes.slides[1].notes_slide.notes_text_frame.text
    without = reopen(build.build_deck(data, analysis, slides, {}, "A Deck", False))
    assert not without.slides[1].has_notes_slide


def test_the_notice_lands_on_every_slide(clean):
    slides = [Slide(id="s1", role="cover", title="A Deck"),
              Slide(id="s2", role="closing", title="Done")]
    written = texts(reopen(deck_of(slides, clean, notice="Built with a model")))
    assert sum(1 for t in written if t == "Built with a model") == 2


def test_date_text_follows_the_language():
    from datetime import date

    assert build.date_text("en", date(2026, 3, 15)) == "Mar 15, 2026"
    assert build.date_text("de", date(2026, 3, 15)) == "15.03.2026"


def test_layout_index_is_kept_once_chosen(clean):
    _, analysis, _ = clean
    s = Slide(id="s1", role="text", layout=analysis.roles["text"][-1])
    assert build.layout_index(analysis, s, {}) == analysis.roles["text"][-1]


def test_layout_rotates_between_the_layouts_of_a_role(clean):
    _, analysis, _ = clean
    counter: dict = {}
    picked = [build.layout_index(analysis, Slide(id=f"s{i}", role="text"), counter) for i in range(4)]
    assert len(set(picked)) == len(analysis.roles["text"])


# Splitting


def test_overflowing_columns_are_split(clean):
    _, analysis, _ = clean
    long = [f"This is point number {n} and it carries a full sentence of its own, twice over, "
            f"so that it certainly wraps" for n in range(1, 9)]
    s = Slide(id="s1", role="two_columns", title="Many points",
              columns=[Column(heading="Left", items=list(long)),
                       Column(heading="Right", items=list(long))],
              notes="First part. Second part. Third part. Fourth part.")
    out = split.split_overflow([s], analysis)
    assert len(out) > 1
    assert sum(len(c.items) for p in out for c in p.columns) == 2 * len(long)
    assert all(p.notes for p in out)
    assert len({p.id for p in out}) == len(out)


def test_a_slide_that_fits_is_not_split(clean):
    _, analysis, _ = clean
    s = Slide(id="s1", role="text", title="Few", body=Body(kind="bullets", items=["One", "Two"]))
    assert split.split_overflow([s], analysis) == [s]


def test_new_id_counts_past_the_highest():
    assert split.new_id({"s1", "s7", "s3"}) == "s8"


def test_notes_are_split_at_a_sentence_end():
    first, second = split.split_notes("One sentence here. Another sentence here.", 0.5)
    assert first == "One sentence here."
    assert second == "Another sentence here."
